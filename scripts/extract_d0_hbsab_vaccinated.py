"""
从「免疫原性原始数据（Cursor整理）-阶段性分析-20260407.xlsx」中按研究编号
提取 62 例有乙肝疫苗接种史受试者的 D0 表面抗体（Anti-HBs）基线数据，
按 5 个组别整理为 Excel，并附组别汇总统计与第 3 册 14.2.3.1.1 FAS 表的交叉验证。

LLOQ 规则（与第 3 册免疫原性分析一致）："<2.00" 按 LLOQ/2 = 1.00 用于 GMC 计算。
"""

import math
import re
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ============ 路径 ============
BASE = Path(r"E:\Cursor Project\2-Scientific-Skills-for-Clinical_Trial")
SRC_XLSX = BASE / "review_materials" / "免疫原性原始数据（Cursor整理）-阶段性分析-20260407.xlsx"
CACHE = BASE / "review_materials" / "_md_cache"
VOL8_MD = CACHE / "第8册_清单1.md"
OUT_XLSX = BASE / "review_materials" / "有接种史_62例_D0表面抗体基线_清单.xlsx"
LLOQ = 2.00
LLOQ_HALF = LLOQ / 2  # = 1.00

# ============ 62 例受试者（从第 8 册 16.2.4.3 + 16.2.4.12 联合确认）============
# 格式：研究编号 -> (组别, 性别, 年龄, 既往接种剂次)
SUBJECTS = {
    # A1 (0,1 月低剂量)
    "112": ("0,1月低剂量组(A1)", "女", 25, "不详"),
    "145": ("0,1月低剂量组(A1)", "男", 42, "三剂"),
    "163": ("0,1月低剂量组(A1)", "女", 47, "二剂"),
    "194": ("0,1月低剂量组(A1)", "女", 47, "三剂"),
    "200": ("0,1月低剂量组(A1)", "女", 26, "三剂"),
    "276": ("0,1月低剂量组(A1)", "女", 22, "三剂"),
    "355": ("0,1月低剂量组(A1)", "男", 22, "三剂"),
    "375": ("0,1月低剂量组(A1)", "男", 18, "三剂"),
    "408": ("0,1月低剂量组(A1)", "男", 34, "不详"),
    "496": ("0,1月低剂量组(A1)", "女", 29, "三剂"),
    "584": ("0,1月低剂量组(A1)", "女", 39, "三剂"),
    "690": ("0,1月低剂量组(A1)", "女", 19, "三剂"),
    "743": ("0,1月低剂量组(A1)", "男", 18, "三剂"),
    # A2 (0,1 月高剂量)
    "124": ("0,1月高剂量组(A2)", "女", 25, "三剂"),
    "219": ("0,1月高剂量组(A2)", "男", 21, "三剂"),
    "345": ("0,1月高剂量组(A2)", "男", 21, "三剂"),
    "369": ("0,1月高剂量组(A2)", "男", 24, "三剂"),
    "378": ("0,1月高剂量组(A2)", "男", 26, "三剂"),
    "390": ("0,1月高剂量组(A2)", "女", 51, "一剂"),
    "412": ("0,1月高剂量组(A2)", "男", 33, "不详"),
    "442": ("0,1月高剂量组(A2)", "男", 43, "一剂"),
    "465": ("0,1月高剂量组(A2)", "女", 54, "不详"),
    "635": ("0,1月高剂量组(A2)", "女", 21, "三剂"),
    "716": ("0,1月高剂量组(A2)", "女", 24, "不详"),
    "727": ("0,1月高剂量组(A2)", "男", 23, "三剂"),
    "728": ("0,1月高剂量组(A2)", "男", 21, "三剂"),
    # B1 (0,2 月低剂量)
    "083": ("0,2月低剂量组(B1)", "女", 34, "不详"),
    "141": ("0,2月低剂量组(B1)", "女", 22, "不详"),
    "169": ("0,2月低剂量组(B1)", "女", 52, "一剂"),
    "173": ("0,2月低剂量组(B1)", "女", 36, "三剂"),
    "218": ("0,2月低剂量组(B1)", "女", 38, "一剂"),
    "270": ("0,2月低剂量组(B1)", "男", 49, "不详"),
    "342": ("0,2月低剂量组(B1)", "男", 19, "三剂"),
    "366": ("0,2月低剂量组(B1)", "女", 23, "三剂"),
    "404": ("0,2月低剂量组(B1)", "女", 42, "三剂"),
    "421": ("0,2月低剂量组(B1)", "女", 40, "三剂"),
    "544": ("0,2月低剂量组(B1)", "男", 45, "三剂"),
    "661": ("0,2月低剂量组(B1)", "男", 33, "一剂"),
    # B2 (0,2 月高剂量)
    "075": ("0,2月高剂量组(B2)", "男", 22, "不详"),
    "298": ("0,2月高剂量组(B2)", "女", 35, "一剂"),
    "347": ("0,2月高剂量组(B2)", "女", 24, "不详"),
    "353": ("0,2月高剂量组(B2)", "男", 25, "三剂"),
    "410": ("0,2月高剂量组(B2)", "男", 26, "三剂"),
    "418": ("0,2月高剂量组(B2)", "男", 44, "不详"),
    "453": ("0,2月高剂量组(B2)", "女", 35, "不详"),
    "500": ("0,2月高剂量组(B2)", "女", 33, "不详"),
    "630": ("0,2月高剂量组(B2)", "男", 38, "三剂"),
    "712": ("0,2月高剂量组(B2)", "女", 49, "三剂"),
    "719": ("0,2月高剂量组(B2)", "女", 31, "不详"),
    # C1 (阳性对照)
    "035": ("阳性对照组(C1)", "男", 27, "不详"),
    "080": ("阳性对照组(C1)", "女", 43, "不详"),
    "110": ("阳性对照组(C1)", "男", 27, "三剂"),
    "170": ("阳性对照组(C1)", "女", 33, "三剂"),
    "181": ("阳性对照组(C1)", "男", 45, "三剂"),
    "240": ("阳性对照组(C1)", "女", 36, "三剂"),
    "405": ("阳性对照组(C1)", "男", 58, "三剂"),
    "416": ("阳性对照组(C1)", "女", 36, "三剂"),
    "438": ("阳性对照组(C1)", "男", 29, "不详"),
    "505": ("阳性对照组(C1)", "女", 41, "二剂"),
    "526": ("阳性对照组(C1)", "男", 39, "二剂"),
    "532": ("阳性对照组(C1)", "女", 37, "不详"),
    "724": ("阳性对照组(C1)", "男", 22, "三剂"),
}

# 第 3 册 14.2.3.1.1 FAS 报告值（FAS 主分析集，交叉验证用）
REPORT_FAS = {
    "0,1月低剂量组(A1)": {"N": 13, "GM": 1.23, "GM_L": 0.90, "GM_U": 1.67, "PosN": 0},
    "0,1月高剂量组(A2)": {"N": 13, "GM": 1.12, "GM_L": 0.95, "GM_U": 1.32, "PosN": 0},
    "0,2月低剂量组(B1)": {"N": 12, "GM": 1.13, "GM_L": 0.86, "GM_U": 1.50, "PosN": 0},
    "0,2月高剂量组(B2)": {"N": 11, "GM": 1.97, "GM_L": 0.69, "GM_U": 5.66, "PosN": 1},
    "阳性对照组(C1)": {"N": 13, "GM": 1.15, "GM_L": 0.94, "GM_U": 1.41, "PosN": 0},
}


# ============ 数据解析 ============
def parse_hbsab(raw) -> tuple[float, str]:
    """解析 Anti-HBs 原始单元格值，返回 (用于GMC的数值, 显示字符串)。
    规则：'＜2.00' / '<2.00' 等低于LLOQ的形式按 LLOQ/2 = 1.00 计。"""
    if raw is None or raw == "":
        return (None, "缺失")
    s = str(raw).strip()
    # 低于 LLOQ 形式
    if s.startswith("<") or s.startswith("＜") or s.startswith("&lt;") or "LLOQ" in s.upper():
        m = re.search(r"(\d+(?:\.\d+)?)", s)
        if m:
            lloq = float(m.group(1))
            return (LLOQ_HALF, f"<{lloq:.2f}")
        return (LLOQ_HALF, f"<{LLOQ:.2f}")
    # 普通数值
    try:
        v = float(s)
        return (v, f"{v:.2f}")
    except ValueError:
        return (None, s)


def load_d0_data() -> dict[str, dict]:
    """从原始 Excel 读取所有 D0 行，返回 {编号: {'hbsab_raw', 'hbsab_val', 'hbsag', ...}}"""
    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
    ws = wb["汇总"]
    d0_map = {}
    for r in range(3, ws.max_row + 1):
        sid = ws.cell(row=r, column=1).value
        if not sid or "-D0" not in str(sid):
            continue
        num = str(sid).split("-")[0].strip().zfill(3)
        hbsab_raw = ws.cell(row=r, column=2).value
        hbsab_val, hbsab_disp = parse_hbsab(hbsab_raw)
        d0_map[num] = {
            "sample_id": sid,
            "hbsag_raw": ws.cell(row=r, column=4).value,
            "hbsab_raw": hbsab_raw,
            "hbsab_val": hbsab_val,
            "hbsab_disp": hbsab_disp,
            "hbcab_raw": ws.cell(row=r, column=6).value,
            "hbeab_raw": ws.cell(row=r, column=8).value,
            "hbeag_raw": ws.cell(row=r, column=10).value,
        }
    return d0_map


# ============ 统计 ============
def gmt_ci(values: list[float], alpha: float = 0.05):
    """对数转换法计算 GMC 与 95% CI。"""
    vals = [v for v in values if v is not None and v > 0]
    n = len(vals)
    if n == 0:
        return (None, None, None, 0)
    log_vals = [math.log(v) for v in vals]
    mean_log = sum(log_vals) / n
    var = sum((x - mean_log) ** 2 for x in log_vals) / (n - 1) if n > 1 else 0
    sd_log = math.sqrt(var)
    gm = math.exp(mean_log)
    se = sd_log / math.sqrt(n) if n > 1 else 0
    z = 1.96  # 95% CI
    gm_l = math.exp(mean_log - z * se)
    gm_u = math.exp(mean_log + z * se)
    return (gm, gm_l, gm_u, n)


# ============ 写 Excel ============
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
GROUP_FILLS = {
    "0,1月低剂量组(A1)": PatternFill("solid", fgColor="DDEBF7"),
    "0,1月高剂量组(A2)": PatternFill("solid", fgColor="DDEBF7"),
    "0,2月低剂量组(B1)": PatternFill("solid", fgColor="FFF2CC"),
    "0,2月高剂量组(B2)": PatternFill("solid", fgColor="FCE4D6"),
    "阳性对照组(C1)": PatternFill("solid", fgColor="E2EFDA"),
}
THIN = Side(border_style="thin", color="B4B4B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_listing_sheet(ws, data_by_group):
    """Sheet 1: 逐例清单"""
    ws.title = "62例_D0表面抗体基线"
    headers = [
        "序号",
        "组别",
        "研究编号",
        "性别",
        "年龄(岁)",
        "D0 样品ID",
        "D0 原始读数(mIU/mL)",
        "用于GMC的数值(mIU/mL)",
        "HBsAg(IU/mL)",
        "Anti-HBc(S/CO)",
        "Anti-HBe(S/CO)",
        "HBeAg(S/CO)",
        "既往接种剂次",
        "免前阳性(≥10 mIU/mL)",
    ]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER

    seq = 0
    row = 2
    for group in [
        "0,1月低剂量组(A1)",
        "0,1月高剂量组(A2)",
        "0,2月低剂量组(B1)",
        "0,2月高剂量组(B2)",
        "阳性对照组(C1)",
    ]:
        for sid in data_by_group[group]:
            seq += 1
            grp, sex, age, doses = SUBJECTS[sid]
            d = data_by_group[group][sid]
            hbsab_val = d["hbsab_val"]
            is_pos = "是" if (hbsab_val is not None and hbsab_val >= 10) else "否"
            row_data = [
                seq,
                group,
                sid,
                sex,
                age,
                d["sample_id"],
                d["hbsab_disp"],
                hbsab_val if hbsab_val is not None else "缺失",
                d["hbsag_raw"],
                d["hbcab_raw"],
                d["hbeab_raw"],
                d["hbeag_raw"],
                doses,
                is_pos,
            ]
            for j, v in enumerate(row_data, 1):
                c = ws.cell(row=row, column=j, value=v)
                c.fill = GROUP_FILLS[group]
                c.border = BORDER
                if j in (1, 2, 3, 4, 5, 6, 14):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            row += 1

    # 列宽
    widths = [6, 18, 10, 6, 8, 14, 18, 18, 14, 14, 14, 14, 12, 14]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"


def write_summary_sheet(ws, data_by_group):
    """Sheet 2: 组别汇总统计 + 交叉验证"""
    ws.title = "组别汇总统计"
    headers = [
        "组别",
        "N",
        "D0 GMC (mIU/mL)",
        "GMC 95% CI 下限",
        "GMC 95% CI 上限",
        "免前阳性 N (≥10 mIU/mL)",
        "免前阳性率",
        "Min",
        "Max",
        "Median",
        "报告 FAS GMC (第3册)",
        "差值 (本次 - 报告)",
    ]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER

    row = 2
    for group in [
        "0,1月低剂量组(A1)",
        "0,1月高剂量组(A2)",
        "0,2月低剂量组(B1)",
        "0,2月高剂量组(B2)",
        "阳性对照组(C1)",
        "低剂量组合计(A1+B1)",
        "高剂量组合计(A2+B2)",
        "全部合计",
    ]:
        if group == "低剂量组合计(A1+B1)":
            subjects = ["0,1月低剂量组(A1)", "0,2月低剂量组(B1)"]
        elif group == "高剂量组合计(A2+B2)":
            subjects = ["0,1月高剂量组(A2)", "0,2月高剂量组(B2)"]
        elif group == "全部合计":
            subjects = [
                "0,1月低剂量组(A1)",
                "0,1月高剂量组(A2)",
                "0,2月低剂量组(B1)",
                "0,2月高剂量组(B2)",
                "阳性对照组(C1)",
            ]
        else:
            subjects = [group]

        all_vals = []
        all_pos = 0
        for sub in subjects:
            for _sid, d in data_by_group[sub].items():
                if d["hbsab_val"] is not None:
                    all_vals.append(d["hbsab_val"])
                    if d["hbsab_val"] >= 10:
                        all_pos += 1

        n = len(all_vals)
        gm, gm_l, gm_u, _ = gmt_ci(all_vals)
        min_v = min(all_vals) if all_vals else None
        max_v = max(all_vals) if all_vals else None
        if all_vals:
            sorted_v = sorted(all_vals)
            if n % 2 == 1:
                median = sorted_v[n // 2]
            else:
                median = (sorted_v[n // 2 - 1] + sorted_v[n // 2]) / 2
        else:
            median = None
        pos_rate = f"{all_pos}/{n} ({all_pos / n * 100:.2f}%)" if n else "—"

        # 报告值
        if group in REPORT_FAS:
            report_gm = REPORT_FAS[group]["GM"]
            report_n = REPORT_FAS[group]["N"]
            diff = (gm - report_gm) if gm is not None else None
            report_disp = f"{report_gm:.2f} (N={report_n})"
            diff_disp = f"{diff:+.3f}" if diff is not None else "—"
        else:
            report_disp = "—（合并组别，无单项报告）"
            diff_disp = "—"

        row_data = [
            group,
            n,
            f"{gm:.3f}" if gm is not None else "—",
            f"{gm_l:.3f}" if gm_l is not None else "—",
            f"{gm_u:.3f}" if gm_u is not None else "—",
            all_pos,
            pos_rate,
            f"{min_v:.2f}" if min_v is not None else "—",
            f"{max_v:.2f}" if max_v is not None else "—",
            f"{median:.2f}" if median is not None else "—",
            report_disp,
            diff_disp,
        ]
        for j, v in enumerate(row_data, 1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if group in GROUP_FILLS:
                c.fill = GROUP_FILLS[group]
            elif "合计" in group:
                c.fill = PatternFill("solid", fgColor="D9D9D9")
                c.font = Font(bold=True)
        row += 1

    # 列宽
    widths = [22, 6, 16, 14, 14, 14, 18, 10, 10, 10, 22, 18]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 36


def write_raw_d0_sheet(ws, data_by_group):
    """Sheet 3: 长格式 D0 原始读数（方便溯源）"""
    ws.title = "D0原始读数(溯源)"
    headers = ["研究编号", "组别", "D0 样品ID", "Anti-HBs 原始读数", "用于GMC数值", "解析说明"]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER

    row = 2
    for group in [
        "0,1月低剂量组(A1)",
        "0,1月高剂量组(A2)",
        "0,2月低剂量组(B1)",
        "0,2月高剂量组(B2)",
        "阳性对照组(C1)",
    ]:
        for sid in data_by_group[group]:
            d = data_by_group[group][sid]
            val = d["hbsab_val"]
            note = ""
            if d["hbsab_disp"].startswith("<"):
                note = f"低于 LLOQ={LLOQ} mIU/mL，按 LLOQ/2={LLOQ_HALF} 计"
            row_data = [
                sid,
                group,
                d["sample_id"],
                d["hbsab_disp"],
                val if val is not None else "缺失",
                note,
            ]
            for j, v in enumerate(row_data, 1):
                c = ws.cell(row=row, column=j, value=v)
                c.border = BORDER
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if group in GROUP_FILLS:
                    c.fill = GROUP_FILLS[group]
            row += 1

    widths = [12, 18, 14, 20, 16, 40]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"


def main():
    print(f"读取原始 Excel: {SRC_XLSX.name}")
    d0_map = load_d0_data()
    print(f"共提取 D0 数据 {len(d0_map)} 例（Excel 中所有受试者）")

    # 关联 62 例
    data_by_group: dict[str, dict] = defaultdict(dict)
    missing = []
    for sid, (group, _sex, _age, _doses) in SUBJECTS.items():
        d = d0_map.get(sid)
        if d is None:
            missing.append(sid)
            continue
        data_by_group[group][sid] = d

    print(
        f"62 例匹配结果: A1={len(data_by_group['0,1月低剂量组(A1)'])}, "
        f"A2={len(data_by_group['0,1月高剂量组(A2)'])}, "
        f"B1={len(data_by_group['0,2月低剂量组(B1)'])}, "
        f"B2={len(data_by_group['0,2月高剂量组(B2)'])}, "
        f"C1={len(data_by_group['阳性对照组(C1)'])}"
    )
    if missing:
        print(f"⚠️ 缺失编号: {missing}")

    # 写 Excel
    wb = openpyxl.Workbook()
    # 删除默认 sheet
    wb.remove(wb.active)
    ws1 = wb.create_sheet()
    ws2 = wb.create_sheet()
    ws3 = wb.create_sheet()

    write_listing_sheet(ws1, data_by_group)
    write_summary_sheet(ws2, data_by_group)
    write_raw_d0_sheet(ws3, data_by_group)

    wb.save(OUT_XLSX)
    print(f"\n已写出: {OUT_XLSX}")


if __name__ == "__main__":
    main()
