"""Dump the full text (including footnotes) of slides 7 and 8 from the main
EOP2 face-to-face expert meeting deck, so every data point and citation marker
on those two pages can be verified against its source.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

PPTX = (
    Path(r"E:\Cursor Project\2-Scientific-Skills-for-Clinical_Trial\review_materials")
    / "TVAX-009 EOP2 面对面专家会议-20260911.pptx"
)
OUT = Path(__file__).with_name("_eop2_slide7_8.txt")

TARGETS = {7, 8}

_buf: list[str] = []


def out(*args) -> None:
    _buf.append(" ".join(str(a) for a in args))


def walk(shapes, depth=0):
    for sh in shapes:
        pad = "  " * (depth + 1)
        if sh.shape_type == 6 or sh.__class__.__name__ == "GroupShape":
            out(f"{pad}[GROUP] {sh.name}")
            walk(sh.shapes, depth + 1)
            continue
        if sh.has_table:
            out(f"{pad}[TABLE] {sh.name}")
            for r in sh.table.rows:
                out(f"{pad}  | " + " | ".join(c.text.replace("\n", " / ") for c in r.cells))
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            out(f"{pad}[TEXT] {sh.text_frame.text.strip()}")
        elif sh.has_text_frame:
            out(f"{pad}[EMPTY-TEXTBOX] {sh.name}")


prs = Presentation(str(PPTX))
out(f"file: {PPTX.name}")
out(f"total slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides, 1):
    if i not in TARGETS:
        continue
    out("=" * 90)
    out(f"SLIDE {i}")
    walk(slide.shapes)
OUT.write_text("\n".join(_buf), encoding="utf-8")
