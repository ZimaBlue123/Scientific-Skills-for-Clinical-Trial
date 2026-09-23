#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
V3：以用户微调版 V2 为基底，仅重建第 13 页流程图（增加细胞免疫采血点通道）。

改动点：
  1) 流程图每条时间轴四类标记：▲ 接种 ｜ ● 体液免疫采血 ｜ ◆ 细胞免疫1组采血 ｜ ◇ 细胞免疫2组采血
  2) 主终点红点+★ 修正到正确时点（旧版 M7；新版 0,1月 M2、0,2月 M3、0,1,6月 M7）
     —— V2 中旧版误将 D0 标为主终点，本次一并修正
  3) 修正 V2 笔误：新版 0,1 月程序 M2 无接种（2 剂在 D0、M1）
  4) 其余页面原样保留用户微调

输入：outputs/TVAX-009_Ⅲ期方案修订要点_20260526vs20260922_V2_backup.pptx（用户微调版）
输出：outputs/TVAX-009_Ⅲ期方案修订要点_20260526vs20260922_V3.pptx
"""
import copy
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

BASE = r"E:\Cursor Project\2-Scientific-Skills-for-Clinical_Trial"
SRC = os.path.join(BASE, "outputs", "TVAX-009_Ⅲ期方案修订要点_20260526vs20260922_V2_backup.pptx")
OUT = os.path.join(BASE, "outputs", "TVAX-009_Ⅲ期方案修订要点_20260526vs20260922_V3.pptx")

FONT = "Arial"
RED = RGBColor(0xC0, 0x00, 0x00)
DARK_RED = RGBColor(0xA1, 0x1C, 0x1D)
PINK = RGBColor(0xE3, 0xB8, 0xB8)
LIGHT_PINK = RGBColor(0xFB, 0xEA, 0xEA)
GRAY = RGBColor(0x59, 0x5F, 0x6B)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _style_run(run, size=10, bold=False, color=BLACK, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls
        latin = parse_xml('<a:latin %s typeface="%s"/>' % (nsdecls("a"), font))
        rPr.append(latin)
    else:
        latin.set("typeface", font)
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = copy.deepcopy(latin)
        ea.tag = qn("a:ea")
        latin.addnext(ea)
    ea.set("typeface", font)
    return run


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return tf


def add_para(tf, text, size=10, bold=False, color=BLACK, first=False, line=None,
             align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    _style_run(r, size, bold, color)
    return p


def rect(slide, left, top, width, height, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, left, top, width, height)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def schedule_card_v3(slide, left, top, width, height, title, rows, head_fill=DARK_RED):
    """四通道流程图卡片。
    rows = [(行标签, [(时点名, vac, hum, cmi1, cmi2, star), ...]), ...]
    标记：▲接种(轴上) ●体液(轴下1) ◆细胞1(轴下2) ◇细胞2(轴下2,空心) 红大圆+★=主终点(替代该点体液圆)
    """
    rect(slide, left, top, width, height, fill=WHITE, line=PINK)
    head_h = Emu(310896)
    rect(slide, left, top, width, head_h, fill=head_fill)
    tf = textbox(slide, left + Emu(109728), top + Emu(48000), width - Emu(219456), head_h)
    add_para(tf, title, size=10.5, bold=True, color=WHITE, first=True)
    axis_l = left + Emu(1250000)
    axis_r = left + width - Emu(140000)
    axis_w = axis_r - axis_l
    row_h = Emu(960000)
    tri_d = dot_d = Emu(100000)
    for ri, (label, pts) in enumerate(rows):
        ry = top + head_h + Emu(40000) + ri * row_h
        lt = textbox(slide, left + Emu(100000), ry, Emu(1120000), Emu(300000))
        add_para(lt, label, size=8.5, bold=True, color=BLACK, first=True)
        axis_y = ry + Emu(400000)
        ln = rect(slide, axis_l, axis_y, axis_w, Emu(12700), fill=PINK)
        n = len(pts)
        span = axis_w - Emu(260000)
        for pi, (name, vac, hum, cmi1, cmi2, star) in enumerate(pts):
            x = axis_l + Emu(130000) + int(span * pi / max(n - 1, 1))
            if vac:
                t = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                           x - tri_d // 2, axis_y - Emu(175000), tri_d, tri_d)
                t.fill.solid(); t.fill.fore_color.rgb = RED
                t.line.fill.background(); t.shadow.inherit = False
            if hum:
                if star:
                    d = Emu(150000)
                    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - d // 2, axis_y + Emu(40000), d, d)
                    c.fill.solid(); c.fill.fore_color.rgb = RED
                    c.line.color.rgb = WHITE; c.line.width = Pt(1); c.shadow.inherit = False
                else:
                    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - dot_d // 2, axis_y + Emu(40000), dot_d, dot_d)
                    c.fill.solid(); c.fill.fore_color.rgb = GRAY
                    c.line.fill.background(); c.shadow.inherit = False
            if cmi1:
                d = Emu(100000)
                off = Emu(-60000) if cmi2 else Emu(0)  # 与◇并排错开
                c = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, x - d // 2 + off, axis_y + Emu(175000), d, d)
                c.fill.solid(); c.fill.fore_color.rgb = DARK_RED
                c.line.fill.background(); c.shadow.inherit = False
            if cmi2:
                d = Emu(100000)
                off = Emu(60000) if cmi1 else Emu(0)  # 与◆并排错开
                c = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, x - d // 2 + off, axis_y + Emu(175000), d, d)
                c.fill.solid(); c.fill.fore_color.rgb = WHITE
                c.line.color.rgb = RED; c.line.width = Pt(1); c.shadow.inherit = False
            lt2 = textbox(slide, x - Emu(280000), axis_y + Emu(310000), Emu(560000), Emu(190000))
            add_para(lt2, name + ("★" if star else ""), size=6.5, bold=star,
                     color=RED if star else GRAY, first=True, align=PP_ALIGN.CENTER)


# 时点数据：(名称, 接种, 体液免疫, 细胞免疫1组, 细胞免疫2组, 主终点)
OLD_ROWS = [
    ("0,1 月程序", [("D0", 1, 1, 0, 0, 0), ("M1", 1, 1, 0, 0, 0), ("M2", 0, 1, 0, 0, 0),
                    ("M6", 0, 1, 0, 0, 0), ("M7", 0, 1, 0, 0, 1), ("M12", 0, 1, 0, 0, 0), ("M24", 0, 1, 0, 0, 0)]),
    ("0,2 月程序", [("D0", 1, 1, 0, 0, 0), ("M1", 0, 1, 0, 0, 0), ("M2", 1, 1, 0, 0, 0), ("M3", 0, 1, 0, 0, 0),
                    ("M6", 0, 1, 0, 0, 0), ("M7", 0, 1, 0, 0, 1), ("M8", 0, 1, 0, 0, 0), ("M12", 0, 1, 0, 0, 0), ("M24", 0, 1, 0, 0, 0)]),
    ("0,1,6 月程序", [("D0", 1, 1, 0, 0, 0), ("M1", 1, 1, 0, 0, 0), ("M6", 1, 1, 0, 0, 0), ("M7", 0, 1, 0, 0, 1),
                      ("M12", 0, 1, 0, 0, 0), ("M18", 0, 1, 0, 0, 0), ("M24", 0, 1, 0, 0, 0)]),
]
NEW_ROWS = [
    ("0,1 月程序", [("D0", 1, 1, 1, 1, 0), ("D1", 0, 0, 0, 1, 0), ("D3", 0, 0, 0, 1, 0), ("D8", 0, 0, 1, 1, 0),
                    ("M1", 1, 1, 1, 0, 0), ("M1+8D", 0, 0, 1, 0, 0), ("M2", 0, 1, 0, 0, 1), ("M3", 0, 1, 0, 0, 0),
                    ("M6", 0, 1, 1, 0, 0), ("M7", 0, 1, 0, 0, 0), ("M13", 0, 1, 0, 0, 0)]),
    ("0,2 月程序", [("D0", 1, 1, 1, 1, 0), ("D1", 0, 0, 0, 1, 0), ("D3", 0, 0, 0, 1, 0), ("D8", 0, 0, 1, 1, 0),
                    ("M1", 0, 1, 0, 0, 0), ("M2", 1, 0, 1, 0, 0), ("M2+8D", 0, 0, 1, 0, 0), ("M3", 0, 1, 0, 0, 1),
                    ("M4", 0, 1, 0, 0, 0), ("M6", 0, 1, 1, 0, 0), ("M8", 0, 1, 0, 0, 0), ("M14", 0, 1, 0, 0, 0)]),
    ("0,1,6 月程序", [("D0", 1, 1, 1, 1, 0), ("D1", 0, 0, 0, 1, 0), ("D3", 0, 0, 0, 1, 0), ("D8", 0, 0, 1, 1, 0),
                      ("M1", 1, 1, 0, 0, 0), ("M6", 1, 1, 1, 0, 0), ("M6+8D", 0, 0, 1, 0, 0), ("M7", 0, 1, 0, 0, 1),
                      ("M8", 0, 1, 0, 0, 0), ("M12", 0, 1, 0, 0, 0), ("M18", 0, 1, 0, 0, 0)]),
]


def main():
    prs = Presentation(SRC)
    slide = prs.slides[12]  # P13
    # 删除旧流程卡与图例（top 介于 2.25M 与 6.3M EMU 之间的形状）；页脚(6419088)与上方表格保留
    removed = 0
    for sp in list(slide.shapes):
        top = sp.top or 0
        if 2250000 <= top < 6300000:
            sp._element.getparent().remove(sp._element)
            removed += 1
    print("removed shapes:", removed)

    cw = Emu(5669280)
    card_top = Emu(2260000)
    schedule_card_v3(slide, M_LEFT := Emu(256032), card_top, cw, Emu(3260000),
                     "旧版 V1.0：接种与体液免疫采血时点（主终点均为 M7，无细胞免疫）",
                     OLD_ROWS, head_fill=GRAY)
    schedule_card_v3(slide, Emu(6236208), card_top, cw, Emu(3260000),
                     "新版 2026-09-22：增设细胞免疫采血点；主终点改为各组“全程接种后 1 个月”",
                     NEW_ROWS, head_fill=RED)
    tf = textbox(slide, Emu(256032), Emu(5620000), Emu(11658600), Emu(700000))
    add_para(tf, "图例：▲ 接种 ｜ ● 体液免疫采血 ｜ ◆ 细胞免疫1组采血（ELISpot＋ICS，100 例） ｜ ◇ 细胞免疫2组采血（pDC 机制层，25 例） ｜ ●（红）＋★ 主要终点采血时点。"
             "末次访视：旧版均为 M24 → 新版 0,1 月 M13、0,2 月 M14、0,1,6 月 M18（全程接种后 12 个月）。细胞免疫 1 组时点：免前、首剂后第 8 天、末剂前、末剂后第 8 天、首剂后 6 个月；2 组时点：免前、D1、D3、D8。",
             size=8, color=GRAY, first=True, line=1.2)

    prs.save(OUT)
    print("saved ->", OUT)


if __name__ == "__main__":
    main()
