#!/usr/bin/env python3
"""Full, per-file extraction of the HEPLISAV-B literature library (PDF + PPTX),
with a yield report so scanned / text-less files cannot be silently missed.

Usage:
    python extract_heplisav_library_full.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader

ROOT = Path(r"E:\Cursor Project\2-Scientific-Skills-for-Clinical_Trial")
BASE = (
    ROOT
    / "review_materials"
    / "文献库-F2F Meeting"
    / "02_同类产品-CpG佐剂与对照疫苗"
    / "HEPLISAV-B Dynavax"
)
OUTDIR = ROOT / "scripts" / "_hep_lib"
MANIFEST = ROOT / "scripts" / "_hep_manifest.txt"


def dump_shape(sh, out, depth=0):
    if sh.shape_type == 6:  # GROUP
        try:
            for sub in sh.shapes:
                dump_shape(sub, out, depth + 1)
        except Exception:  # noqa: BLE001
            pass
        return
    if getattr(sh, "has_table", False) and sh.has_table:
        for row in sh.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            dedup = []
            for c in cells:
                if not dedup or dedup[-1] != c:
                    dedup.append(c)
            out.append(" | ".join(dedup))
        return
    if sh.has_text_frame:
        t = sh.text_frame.text.strip()
        if t:
            out.append(t)


def read_pptx(path: Path) -> tuple[str, int]:
    prs = Presentation(str(path))
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"\n--- slide {i} ---")
        for sh in slide.shapes:
            dump_shape(sh, lines)
    return "\n".join(lines), len(prs.slides)


def read_pdf(path: Path) -> tuple[str, int]:
    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        try:
            txt = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            txt = ""
        parts.append(f"\n--- p.{i} ---\n{txt}")
    return "\n".join(parts), len(reader.pages)


def main() -> int:
    OUTDIR.mkdir(parents=True, exist_ok=True)
    files = sorted([p for p in BASE.rglob("*") if p.suffix.lower() in (".pdf", ".pptx")])
    rows = []
    for f in files:
        rel = f.relative_to(BASE)
        try:
            if f.suffix.lower() == ".pdf":
                text, n = read_pdf(f)
                kind = "pdf"
            else:
                text, n = read_pptx(f)
                kind = "pptx"
        except Exception as exc:  # noqa: BLE001
            rows.append(f"[ERROR] {rel}\t{exc}")
            continue
        slug = str(rel).replace("\\", "__").replace("/", "__")
        (OUTDIR / f"{slug}.txt").write_text(text, encoding="utf-8")
        rows.append(f"{len(text)}\t{kind}\t{n}\t{rel}")

    MANIFEST.write_text("\n".join(rows), encoding="utf-8")
    print(f"files={len(files)} -> {OUTDIR}")
    print("\n".join(rows))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
