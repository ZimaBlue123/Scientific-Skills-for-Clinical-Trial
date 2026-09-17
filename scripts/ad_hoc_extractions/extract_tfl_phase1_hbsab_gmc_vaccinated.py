"""
从 TVAX-009（重组乙型肝炎疫苗，汉逊酵母，CpG和铝佐剂）Ⅰ期统计分析报告
第9册（清单1，DOCX）中提取「有乙肝疫苗接种史」人群各时间点的乙肝表面抗体(anti-HBs) GMC。

数据源（第9册 = 清单1，FAS，直接读 DOCX 表格）：
  1. 表16.2.5.9  既往乙型肝炎疫苗接种史清单        —— "有接种史"受试者名单
  2. 表16.2.5.10 既往加强免疫乙型肝炎疫苗接种史清单 —— 加强免疫史（本试验为"无相关数据"）
  3. 表16.2.6.1  体液免疫结果清单                  —— 逐例 5 个时间点 anti-HBs(mIU/mL)

时间点（5 个）：免前 / 第1剂免后1个月 / 第2剂免后1个月 / 第3剂接种前 / 第3剂免后1个月

LLOQ 规则：anti-HBs "<2.00" 按 LLOQ/2 = 1.00 计（与报告免疫原性分析一致）。
注：表16.2.6.1 的 anti-HBs 为中科院「乙肝两对半检测」定量值(mIU/mL)，与免疫原性专项
检测的 GMC 是两套数据；本清单统一采用表16.2.6.1 值以保证 5 个时间点口径一致。

输出：review_materials/有接种史人群_各时间点乙肝表面抗体GMC_清单.xlsx
  Sheet1 = 逐例清单（受试者 × 5 时间点 anti-HBs 原始值 + 用于GMC的数值 + 接种史信息）
  Sheet2 = GMC 汇总（4 组别 × 5 时间点：N / GMC / 95%CI / 阳性N / 阳性率）
"""

from __future__ import annotations

import math
import re
from pathlib import Path

import openpyxl
from docx import Document
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(r"E:\Cursor Project\2-Scientific-Skills-for-Clinical_Trial")
SRC_DOCX = (
    BASE
    / "review_materials"
    / "TFL-Phase1"
    / "远大赛威信-重组乙型肝炎疫苗（汉逊酵母，CpG和铝佐剂）Ⅰ期统计分析报告-第9册（清单1）.docx"
)
OUT_XLSX = BASE / "review_materials" / "有接种史人群_各时间点乙肝表面抗体GMC_清单.xlsx"

LLOQ = 2.00
LLOQ_HALF = LLOQ / 2  # = 1.00
POS_THRESHOLD = 10.0  # 阳性判定阈值 (mIU/mL)

GROUPS = ["低剂量组", "中剂量组", "高剂量组", "对照组"]
ID_RE = re.compile(r"^[DGZ]\d{2}$")

TIMEPOINTS = ["免前", "第1剂免后1个月", "第2剂免后1个月", "第3剂接种前", "第3剂免后1个月"]


def cell_text(cell) -> str:
    return cell.text.strip()


def row_texts(row) -> list[str]:
    return [cell_text(c) for c in row.cells]


# ============ 定位表格 ============
def find_tables(doc):
    """按表头特征定位：接种史表、加强免疫表、体液免疫表。返回 (vacc_table, booster_table, immuno_table)。"""
    vacc = booster = immuno = None
    for t in doc.tables:
        if len(t.rows) < 1:
            continue
        hdr = row_texts(t.rows[0])
        joined = " ".join(hdr)
        # 体液免疫结果清单：首行含"免前乙肝两对半检测结果"且含"第3剂免后"
        if "免前乙肝两对半检测结果" in joined and "第3剂免后" in joined:
            immuno = t
        # 接种史：首行含"疫苗名称"且含"首次免疫"
        elif "疫苗名称" in joined and "首次免疫" in joined and "加强免疫" not in joined:
            vacc = t
        elif "疫苗名称" in joined and "加强免疫" in joined:
            booster = t
    return vacc, booster, immuno


# ============ 解析接种史 ============
def extract_vaccination_history(vacc_table, booster_table) -> list[dict]:
    """解析表16.2.5.9（+16.2.5.10）接种史。列：0组别 1年龄组 2编号 3年龄 4性别 6疫苗名称 7剂次 9日期 12地点。"""
    rows = []

    def collect(table, src):
        for r in table.rows[1:]:
            cells = row_texts(r)
            if len(cells) < 8:
                continue
            if not ID_RE.match(cells[2]):
                continue
            if "无相关数据" in cells:
                continue
            rows.append(
                {
                    "group": cells[0],
                    "age_group": cells[1],
                    "sid": cells[2],
                    "age": cells[3],
                    "sex": cells[4],
                    "vaccine": cells[6] if len(cells) > 6 else "",
                    "doses": cells[7] if len(cells) > 7 else "",
                    "date": cells[9] if len(cells) > 9 else "",
                    "site": cells[12] if len(cells) > 12 else "",
                    "src": src,
                }
            )

    collect(vacc_table, "既往乙肝疫苗接种史")
    if booster_table is not None:
        collect(booster_table, "既往加强免疫乙型肝炎疫苗接种史")
    return rows


# ============ 解析体液免疫结果清单 ============
def extract_humoral_immuno(immuno_table) -> dict:
    """解析表16.2.6.1 体液免疫结果清单(FAS)。
    数据行：0组别 1年龄组 2编号 3年龄 4性别，anti-HBs 位于列 6/11/13/15/17（0-based）。"""
    records = {}
    hbsab_cols = [6, 11, 13, 15, 17]
    for r in immuno_table.rows[2:]:  # 跳过两行表头
        cells = row_texts(r)
        if len(cells) < 18:
            continue
        sid = cells[2]
        if not ID_RE.match(sid):
            continue
        records[sid] = {
            "group": cells[0],
            "age_group": cells[1],
            "age": cells[3],
            "sex": cells[4],
            "tp_raw": [cells[c] if c < len(cells) else "" for c in hbsab_cols],
        }
    return records


def parse_hbsab(raw: str):
    """解析 anti-HBs 原始值。'<2.00' 等按 LLOQ/2=1.00；空/NA 返回 (None, 原始)。"""
    if raw is None:
        return (None, "缺失")
    s = str(raw).strip()
    if s in ("", "NA", "-", "无"):
        return (None, "缺失")
    if s.startswith("<") or s.startswith("＜"):
        return (LLOQ_HALF, s)
    try:
        v = float(s)
        return (v, s)
    except ValueError:
        return (None, s)


# ============ 统计 ============
def gmc_ci(values: list[float]):
    """对数转换法计算 GMC 与 95% CI。返回 (N, GMC, CI下, CI上)。"""
    vals = [v for v in values if v is not None and v > 0]
    n = len(vals)
    if n == 0:
        return (0, None, None, None)
    log_vals = [math.log(v) for v in vals]
    mean_log = sum(log_vals) / n
    if n > 1:
        var = sum((x - mean_log) ** 2 for x in log_vals) / (n - 1)
        se = math.sqrt(var) / math.sqrt(n)
    else:
        se = 0.0
    gm = math.exp(mean_log)
    z = 1.96
    return (n, gm, math.exp(mean_log - z * se), math.exp(mean_log + z * se))


# ============ 写 Excel ============
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
GROUP_FILLS = {
    "低剂量组": PatternFill("solid", fgColor="DDEBF7"),
    "中剂量组": PatternFill("solid", fgColor="FFF2CC"),
    "高剂量组": PatternFill("solid", fgColor="FCE4D6"),
    "对照组": PatternFill("solid", fgColor="E2EFDA"),
}
TOTAL_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(border_style="thin", color="B4B4B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_listing_sheet(ws, subjects):
    ws.title = "逐例清单"
    headers = [
        "序号",
        "组别",
        "年龄组",
        "研究编号",
        "年龄(岁)",
        "性别",
        "疫苗名称",
        "首次免疫已完成剂次",
        "首剂接种日期",
        "接种地点",
        "接种史来源",
        *[f"{tp}\n原始值(mIU/mL)" for tp in TIMEPOINTS],
        *[f"{tp}\n用于GMC(mIU/mL)" for tp in TIMEPOINTS],
    ]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER

    seq = 0
    row = 2
    for subj in subjects:
        seq += 1
        base = [
            seq,
            subj["group"],
            subj["age_group"],
            subj["sid"],
            subj["age"],
            subj["sex"],
            subj["vaccine"],
            subj["doses"],
            subj["date"],
            subj["site"],
            subj["src"],
        ]
        raw_disp = [d[1] for d in subj["tp"]]
        val_disp = [d[0] if d[0] is not None else "缺失" for d in subj["tp"]]
        row_data = base + raw_disp + val_disp
        for j, v in enumerate(row_data, 1):
            c = ws.cell(row=row, column=j, value=v)
            c.fill = GROUP_FILLS.get(subj["group"], PatternFill())
            c.border = BORDER
            c.alignment = CENTER
        row += 1

    widths = [5, 9, 9, 9, 7, 6, 13, 12, 11, 16, 14] + [11] * 5 + [11] * 5
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "G2"


def write_summary_sheet(ws, summary, order):
    ws.title = "GMC汇总"
    headers = [
        "组别",
        "时间点",
        "N",
        "GMC(mIU/mL)",
        "95%CI下限",
        "95%CI上限",
        "阳性N(≥10)",
        "阳性率",
    ]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER

    row = 2
    for group in order:
        for i, tp in enumerate(TIMEPOINTS):
            s = summary[group][i]
            n, gm, gm_l, gm_u = s["gmc"]
            pos_n = s["pos_n"]
            pos_rate = f"{pos_n}/{n} ({pos_n / n * 100:.2f}%)" if n else "—"
            row_data = [
                group if i == 0 else "",
                tp,
                n,
                f"{gm:.2f}" if gm is not None else "—",
                f"{gm_l:.2f}" if gm_l is not None else "—",
                f"{gm_u:.2f}" if gm_u is not None else "—",
                pos_n,
                pos_rate,
            ]
            for j, v in enumerate(row_data, 1):
                c = ws.cell(row=row, column=j, value=v)
                c.fill = GROUP_FILLS.get(group, PatternFill())
                c.border = BORDER
                c.alignment = CENTER
            row += 1

    # 全体合计
    for i, tp in enumerate(TIMEPOINTS):
        merged_vals = []
        for group in order:
            merged_vals.extend(summary[group][i]["vals"])
        n2, gm2, gm_l2, gm_u2 = gmc_ci(merged_vals)
        pos2 = sum(1 for v in merged_vals if v >= POS_THRESHOLD)
        pos_rate2 = f"{pos2}/{n2} ({pos2 / n2 * 100:.2f}%)" if n2 else "—"
        row_data = [
            "全体合计" if i == 0 else "",
            tp,
            n2,
            f"{gm2:.2f}" if gm2 is not None else "—",
            f"{gm_l2:.2f}" if gm_l2 is not None else "—",
            f"{gm_u2:.2f}" if gm_u2 is not None else "—",
            pos2,
            pos_rate2,
        ]
        for j, v in enumerate(row_data, 1):
            c = ws.cell(row=row, column=j, value=v)
            c.fill = TOTAL_FILL
            c.border = BORDER
            c.alignment = CENTER
            c.font = Font(bold=True)
        row += 1

    widths = [10, 16, 6, 13, 12, 12, 12, 15]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def main():
    doc = Document(str(SRC_DOCX))
    vacc_table, booster_table, immuno_table = find_tables(doc)
    assert vacc_table is not None, "未找到接种史表"
    assert immuno_table is not None, "未找到体液免疫结果清单表"

    # 1. 接种史
    history = extract_vaccination_history(vacc_table, booster_table)
    print(f"[接种史] 表16.2.5.9(+16.2.5.10) 解析到 {len(history)} 例有乙肝疫苗接种史")
    from collections import Counter

    print("    组别分布:", dict(Counter(r["group"] for r in history)))

    # 2. 体液免疫
    immuno = extract_humoral_immuno(immuno_table)
    print(f"[体液免疫] 表16.2.6.1 解析到 {len(immuno)} 例（全体FAS）")

    # 3. 交叉匹配
    subjects = []
    missing = []
    seen = set()
    for r in history:
        if r["sid"] in seen:
            continue  # 同一受试者若同时出现在两张接种史表，只取一次
        seen.add(r["sid"])
        rec = immuno.get(r["sid"])
        if rec is None:
            missing.append(r["sid"])
            continue
        tp = [parse_hbsab(raw) for raw in rec["tp_raw"]]
        subjects.append(
            {
                "group": r["group"],
                "age_group": r["age_group"],
                "sid": r["sid"],
                "age": r["age"],
                "sex": r["sex"],
                "vaccine": r["vaccine"],
                "doses": r["doses"],
                "date": r["date"],
                "site": r["site"],
                "src": r["src"],
                "tp": tp,
            }
        )
    print(f"[匹配] {len(subjects)} 例有接种史受试者匹配到体液免疫数据")
    if missing:
        print(f"    ⚠️ 未匹配编号: {missing}")

    group_order_map = {g: i for i, g in enumerate(GROUPS)}
    subjects.sort(key=lambda s: (group_order_map[s["group"]], s["age_group"], s["sid"]))

    # 4. GMC 汇总
    summary = {}
    for g in GROUPS:
        summary[g] = []
        g_subjects = [s for s in subjects if s["group"] == g]
        for i in range(len(TIMEPOINTS)):
            vals = [s["tp"][i][0] for s in g_subjects if s["tp"][i][0] is not None]
            n, gm, gm_l, gm_u = gmc_ci(vals)
            pos_n = sum(1 for v in vals if v >= POS_THRESHOLD)
            summary[g].append({"gmc": (n, gm, gm_l, gm_u), "pos_n": pos_n, "vals": vals})

    print("\n===== GMC 汇总核对 =====")
    for g in GROUPS:
        print(f"\n【{g}】N={sum(1 for s in subjects if s['group'] == g)}")
        for i, tp in enumerate(TIMEPOINTS):
            n, gm, gm_l, gm_u = summary[g][i]["gmc"]
            pos_n = summary[g][i]["pos_n"]
            gmc_s = f"{gm:.2f} ({gm_l:.2f}-{gm_u:.2f})" if gm is not None else "—"
            print(f"    {tp:<10} N={n:<3} GMC={gmc_s:<22} 阳性={pos_n}")

    # 5. 写 Excel
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws1 = wb.create_sheet()
    ws2 = wb.create_sheet()
    write_listing_sheet(ws1, subjects)
    write_summary_sheet(ws2, summary, GROUPS)
    wb.save(OUT_XLSX)
    print(f"\n已写出: {OUT_XLSX}")


if __name__ == "__main__":
    main()
