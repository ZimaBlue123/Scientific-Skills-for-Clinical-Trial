"""通篇导出 PPT 全部内容（正文 / 表格 / 组合内形状 / 备注），供审阅数据矛盾与表述错误。"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

SRC = Path(
    r"E:\Cursor Project\2-Scientific-Skills-for-Clinical_Trial\review_materials"
    r"\TVAX-009 EOP2 面对面专家会议-20260915（含逐字稿备注）.pptx"
)
OUT = Path(__file__).with_name("_deck_0915_full.txt")

_buf: list[str] = []


def w(s: str = "") -> None:
    _buf.append(s)


def walk(shapes, depth: int = 0) -> None:
    pad = "  " * (depth + 2)
    for sh in shapes:
        cls = sh.__class__.__name__
        if cls == "GroupShape" or sh.shape_type == 6:
            w(f"{pad}[GROUP] {sh.name}")
            walk(sh.shapes, depth + 1)
            continue
        if getattr(sh, "has_table", False):
            tbl = sh.table
            w(f"{pad}[TABLE] {sh.name}")
            for _r_i, row in enumerate(tbl.rows):
                cells = [c.text.replace("\n", " / ").strip() for c in row.cells]
                w(f"{pad}  | " + " | ".join(cells))
            continue
        if sh.has_text_frame:
            txt = sh.text_frame.text.strip()
            if txt:
                for line in txt.split("\n"):
                    line = line.strip()
                    if line:
                        w(f"{pad}[T] {line}")
        if sh.has_chart if hasattr(sh, "has_chart") else False:
            w(f"{pad}[CHART] {sh.name}")


def main() -> int:
    if not SRC.exists():
        w(f"NOT FOUND: {SRC}")
        OUT.write_text("\n".join(_buf), encoding="utf-8")
        return 1
    prs = Presentation(str(SRC))
    w(f"file: {SRC.name}")
    w(f"slides: {len(prs.slides)}")
    w(f"size: {Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in")
    w("")
    for i, slide in enumerate(prs.slides, 1):
        w("=" * 90)
        w(f"SLIDE {i}")
        w("=" * 90)
        walk(slide.shapes, 0)
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame.text.strip()
            if nt:
                w("  --- NOTES ---")
                for line in nt.split("\n"):
                    line = line.rstrip()
                    if line.strip():
                        w(f"  [N] {line}")
        w("")
    OUT.write_text("\n".join(_buf), encoding="utf-8")
    print(f"written {OUT} ({len(_buf)} lines)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
