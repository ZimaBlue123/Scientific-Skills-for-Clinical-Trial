#!/usr/bin/env python3
"""Dump raw XML of selected shapes on given slides (to inspect colors/fills).

Usage:
    python dump_slide_xml.py <pptx> <out.txt> <slide> [shape_name_substr ...]
"""

from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation


def main(argv: list[str]) -> int:
    src = Path(argv[1])
    out_path = Path(argv[2])
    slide_no = int(argv[3])
    names = argv[4:]

    prs = Presentation(str(src))
    slide = prs.slides[slide_no - 1]

    out = []
    for sh in slide.shapes:
        if names and not any(n in sh.name for n in names):
            continue
        out.append(f"===== shape {sh.name!r} ({sh.shape_type}) =====")
        out.append(etree.tostring(sh._element, pretty_print=True).decode("utf-8"))
    # also dump table XML for tables
    for sh in slide.shapes:
        if getattr(sh, "has_table", False) and sh.has_table:
            if names and not any(n in sh.name for n in names):
                continue
            out.append(f"===== TABLE {sh.name!r} first rows =====")
            tbl = sh.table
            for ri in range(min(2, len(tbl.rows))):
                cel = tbl.rows[ri].cells[0]._tc
                out.append(etree.tostring(cel, pretty_print=True).decode("utf-8")[:4000])
    out_path.write_text("\n".join(out), encoding="utf-8")
    print(f"wrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
