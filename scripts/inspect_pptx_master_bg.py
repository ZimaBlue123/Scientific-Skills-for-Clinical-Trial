#!/usr/bin/env python3
"""Dump slide-master shapes and theme colour scheme of a reference PPTX.

Usage:
    python inspect_pptx_master_bg.py <pptx> <out.txt>
"""

from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def inch(v):
    try:
        return f"{Emu(int(v)).inches:.2f}in"
    except Exception:  # noqa: BLE001
        return str(v)


def main(argv):
    src = Path(argv[1])
    out_path = Path(argv[2])
    prs = Presentation(str(src))
    out = []

    # ---- theme colours -------------------------------------------------
    out.append("===== THEME =====")
    for part in prs.part.package.iter_parts():
        pn = str(part.partname)
        if pn.startswith("/ppt/theme/theme"):
            root = etree.fromstring(part.blob)
            for tag in (
                "dk1",
                "lt1",
                "dk2",
                "lt2",
                "accent1",
                "accent2",
                "accent3",
                "accent4",
                "accent5",
                "accent6",
                "hlink",
                "folHlink",
            ):
                for n in root.findall(f".//{{{A}}}clrScheme/{{{A}}}{tag}/*"):
                    out.append(f"  {tag}: {n.tag.split('}')[-1]} = {n.attrib}")
            for ftype in ("majorFont", "minorFont"):
                for n in root.findall(f".//{{{A}}}fontScheme/{{{A}}}{ftype}/{{{A}}}latin"):
                    out.append(f"  {ftype}/latin: {n.attrib}")
            out.append(f"  (theme part: {pn})")
            out.append("")
            break

    # ---- master shapes -------------------------------------------------
    for mi, master in enumerate(prs.slide_masters):
        out.append(f"===== MASTER {mi} shapes =====")
        for sh in master.shapes:
            out.append(
                f"  - {sh.name!r} type={sh.shape_type} "
                f"L={inch(sh.left)} T={inch(sh.top)} W={inch(sh.width)} H={inch(sh.height)}"
            )
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(f"      text={sh.text_frame.text[:100]!r}")
            try:
                if sh.fill.type == 1:
                    out.append(f"      fill=#{sh.fill.fore_color.rgb}")
            except Exception:  # noqa: BLE001
                pass
        # raw bg element
        bg = master.element.find(f"{{{P}}}bg")
        if bg is not None:
            out.append("  master bg XML:")
            out.append("   " + etree.tostring(bg, pretty_print=True).decode()[:1200])
        out.append("")

    # ---- layout 空白 / 仅标题 shapes -------------------------------------
    for mi, master in enumerate(prs.slide_masters):
        for layout in master.slide_layouts:
            if layout.name in ("空白", "仅标题", "标题和内容", "图文內容页一"):
                out.append(f"===== MASTER {mi} LAYOUT {layout.name!r} shapes =====")
                for sh in layout.shapes:
                    out.append(
                        f"  - {sh.name!r} type={sh.shape_type} "
                        f"L={inch(sh.left)} T={inch(sh.top)} "
                        f"W={inch(sh.width)} H={inch(sh.height)}"
                    )
                    if sh.has_text_frame and sh.text_frame.text.strip():
                        out.append(f"      text={sh.text_frame.text[:80]!r}")
                    try:
                        if sh.fill.type == 1:
                            out.append(f"      fill=#{sh.fill.fore_color.rgb}")
                    except Exception:  # noqa: BLE001
                        pass
                for ph in layout.placeholders:
                    out.append(
                        f"    PH idx={ph.placeholder_format.idx} "
                        f"type={ph.placeholder_format.type} name={ph.name!r} "
                        f"L={inch(ph.left)} T={inch(ph.top)} "
                        f"W={inch(ph.width)} H={inch(ph.height)}"
                    )
                out.append("")

    out_path.write_text("\n".join(out), encoding="utf-8")
    print(f"wrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
