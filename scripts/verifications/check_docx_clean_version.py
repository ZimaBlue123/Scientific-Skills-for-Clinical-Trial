#!/usr/bin/env python3
"""
check_docx_clean_version.py
Check Word (.docx) files for "clean version" compliance.

Flags:
  (a) Red font colour
  (b) Blue font colour (hyperlinks excluded)
  (c) Yellow highlight background
  (d) Comments / annotations

Usage:
    python scripts/check_docx_clean_version.py <folder_or_file> [--json]
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

# ── colour constants ────────────────────────────────────────────────────────
RED_RGBS = {
    "FF0000",
    "ff0000",
    "CC0000",
    "cc0000",
    "C00000",
    "c00000",
    "DD0000",
    "dd0000",
    "EE0000",
    "ee0000",
}

BLUE_RGBS = {
    "0000FF",
    "0000ff",
    "0563C1",
    "0563c1",
    "0070C0",
    "0070c0",
    "4472C4",
    "4472c4",
    "1F4E79",
    "1f4e79",
    "2E75B6",
    "2e75b6",
    "002060",
    "0000CC",
    "0000cc",
    "0000EE",
    "0000ee",
    "4F81BD",
    "4f81bd",
    "1155CC",
    "1155cc",
    "2F5496",
    "2f5496",
    "5B9BD5",
    "5b9bd5",
}

YELLOW_HIGHLIGHTS = {"yellow", "7"}

YELLOW_FILL_RGBS = {
    "FFFF00",
    "ffff00",
    "FFFF99",
    "ffff99",
    "FFF200",
    "fff200",
    "FFD700",
    "ffd700",
    "FFC000",
    "ffc000",
    "FFE599",
    "ffe599",
    "FEFF00",
    "feff00",
}


def _is_colour_close(hex_val: str, target_set: set[str]) -> bool:
    if not hex_val:
        return False
    cleaned = hex_val.strip().lstrip("#").upper()
    return cleaned in {v.upper() for v in target_set}


@dataclass
class Issue:
    issue_type: str
    location: str
    snippet: str


@dataclass
class FileResult:
    filepath: str
    filename: str
    is_clean: bool
    issues: list[Issue] = field(default_factory=list)


def _get_run_colour(run) -> str | None:
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        return None
    color_el = rPr.find(qn("w:color"))
    if color_el is None:
        return None
    val = color_el.get(qn("w:val"))
    if val and val.lower() != "auto":
        return val
    return None


def _get_run_highlight(run) -> str | None:
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        return None
    hl_el = rPr.find(qn("w:highlight"))
    if hl_el is None:
        return None
    return hl_el.get(qn("w:val"))


def _run_is_in_hyperlink(run) -> bool:
    parent = run._element.getparent()
    while parent is not None:
        if parent.tag == qn("w:hyperlink"):
            return True
        parent = parent.getparent()
    return False


def _get_run_shading_fill(run) -> str | None:
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        return None
    shd_el = rPr.find(qn("w:shd"))
    if shd_el is None:
        return None
    fill = shd_el.get(qn("w:fill"))
    if fill and fill.lower() not in ("auto", "ffffff", "none"):
        return fill
    return None


def _check_runs_in_paragraphs(paragraphs, loc_prefix, result):
    """Shared logic for checking runs in a list of paragraphs."""
    for p_idx, para in enumerate(paragraphs, start=1):
        for run in para.runs:
            text_snippet = run.text.strip()
            if not text_snippet:
                continue

            snippet = text_snippet[:60]
            loc = f"{loc_prefix} Paragraph {p_idx}" if loc_prefix else f"Paragraph {p_idx}"

            colour = _get_run_colour(run)
            if colour and _is_colour_close(colour, RED_RGBS):
                result.issues.append(Issue("red_font", loc, snippet))

            if colour and _is_colour_close(colour, BLUE_RGBS):
                if not _run_is_in_hyperlink(run):
                    result.issues.append(Issue("blue_font", loc, snippet))

            hl = _get_run_highlight(run)
            if hl and hl.lower() in YELLOW_HIGHLIGHTS:
                result.issues.append(Issue("yellow_highlight", loc, snippet))

            shd_fill = _get_run_shading_fill(run)
            if shd_fill and _is_colour_close(shd_fill, YELLOW_FILL_RGBS):
                result.issues.append(Issue("yellow_highlight_shading", loc, snippet))


def check_docx(filepath: Path) -> FileResult:
    result = FileResult(
        filepath=str(filepath),
        filename=filepath.name,
        is_clean=True,
    )

    try:
        doc = Document(str(filepath))
    except Exception as exc:
        result.issues.append(Issue("error", "N/A", f"Cannot open file: {exc}"))
        result.is_clean = False
        return result

    # 1. Check body paragraphs
    _check_runs_in_paragraphs(doc.paragraphs, "", result)

    # 2. Check tables
    for t_idx, table in enumerate(doc.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            for c_idx, cell in enumerate(row.cells, start=1):
                loc_prefix = f"Table {t_idx} Row {r_idx} Cell {c_idx}"
                _check_runs_in_paragraphs(cell.paragraphs, loc_prefix, result)

    # 3. Check headers and footers
    for section in doc.sections:
        for hf_name, hf in [
            ("Header", section.header),
            ("Footer", section.footer),
        ]:
            if hf and hf.is_linked_to_previous is False:
                _check_runs_in_paragraphs(hf.paragraphs, hf_name, result)

    # 4. Check comments
    import zipfile

    try:
        with zipfile.ZipFile(str(filepath), "r") as zf:
            if "word/comments.xml" in zf.namelist():
                comments_xml = zf.read("word/comments.xml")
                tree = etree.fromstring(comments_xml)
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                comments = tree.findall(".//w:comment", ns)
                for i, comment in enumerate(comments, start=1):
                    author = comment.get(qn("w:author"), "Unknown")
                    texts = comment.findall(".//w:t", ns)
                    comment_text = "".join(t.text or "" for t in texts).strip()
                    snippet = comment_text[:80] if comment_text else "(empty)"
                    result.issues.append(Issue("comment", f"Comment {i} by {author}", snippet))
    except Exception:
        pass

    result.is_clean = len(result.issues) == 0
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Check .docx files for clean-version compliance.",
    )
    parser.add_argument("target", help="File or folder to check")
    parser.add_argument("--json", action="store_true", help="Output as JSON")
    args = parser.parse_args(argv)

    target = Path(args.target)
    if not target.exists():
        print(f"ERROR: path not found: {target}", file=sys.stderr)
        return 1

    files: list[Path] = []
    if target.is_file():
        files.append(target)
    else:
        for root, _dirs, fnames in os.walk(target):
            for fn in sorted(fnames):
                if fn.lower().endswith(".docx") and not fn.startswith("~$"):
                    files.append(Path(root) / fn)

    if not files:
        print("No .docx files found.", file=sys.stderr)
        return 1

    results: list[FileResult] = []
    for f in files:
        print(f"Checking: {f.name} ...", file=sys.stderr)
        results.append(check_docx(f))

    if args.json:
        print(json.dumps([asdict(r) for r in results], ensure_ascii=False, indent=2))
    else:
        for r in results:
            status = "✅ CLEAN" if r.is_clean else "❌ NOT CLEAN"
            print(f"\n{'=' * 78}")
            print(f"  {status}  {r.filename}")
            print(f"  Path: {r.filepath}")
            if r.issues:
                print(f"  Issues found: {len(r.issues)}")
                for issue in r.issues:
                    print(f"    [{issue.issue_type}] {issue.location}")
                    print(f'      "{issue.snippet}"')
            print(f"{'=' * 78}")

        clean_count = sum(1 for r in results if r.is_clean)
        dirty_count = len(results) - clean_count
        print(
            f"\n--- Summary: {len(results)} files checked, "
            f"{clean_count} clean, {dirty_count} not clean ---"
        )

    return 0


if __name__ == "__main__":
    sys.exit(main())
