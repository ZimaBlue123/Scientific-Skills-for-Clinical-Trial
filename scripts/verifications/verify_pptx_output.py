"""
PPTX 输出校验工具：
  1. 打印每页的形状几何与文本（compact）
  2. 标记超出画布边界的形状（新增页排版体检）
用法:
  python scripts/verify_pptx_output.py <pptx路径> <输出txt> [页号 ...]
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

EMU_IN = 914400


def main() -> int:
    if len(sys.argv) < 3:
        print(__doc__)
        return 1
    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    want = [int(x) for x in sys.argv[3:]]

    prs = Presentation(str(src))
    W = prs.slide_width / EMU_IN
    H = prs.slide_height / EMU_IN

    lines = [f"slide size: {W:.3f} x {H:.3f} in", f"slides: {len(list(prs.slides))}", ""]
    for i, slide in enumerate(prs.slides, start=1):
        if want and i not in want:
            continue
        lines.append("=" * 12 + f" SLIDE {i} (layout={slide.slide_layout.name}) " + "=" * 12)
        for shp in slide.shapes:
            try:
                L, T = shp.left / EMU_IN, shp.top / EMU_IN
                w, h = shp.width / EMU_IN, shp.height / EMU_IN
            except TypeError:
                L = T = w = h = float("nan")
            flag = ""
            if L == L and (L < -0.01 or T < -0.01 or L + w > W + 0.01 or T + h > H + 0.01):
                flag = "  <<< OUT OF CANVAS"
            lines.append(
                f"  - {shp.name:<16} L={L:6.3f} T={T:6.3f} W={w:6.3f} H={h:6.3f}"
                f"  R={L + w:6.3f} B={T + h:6.3f}{flag}"
            )
            if shp.has_text_frame and shp.text_frame.text.strip():
                for para in shp.text_frame.paragraphs:
                    if para.text.strip():
                        lines.append(f"      | {para.text}")
            if shp.has_table:
                tbl = shp.table
                lines.append(f"      [TABLE {len(tbl.rows)}x{len(tbl.columns)}]")
                for r in tbl.rows:
                    lines.append(
                        "      | " + " || ".join(c.text.replace("\n", " ") for c in r.cells)
                    )
        lines.append("")

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"[OK] {out}  ({len(lines)} lines)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
