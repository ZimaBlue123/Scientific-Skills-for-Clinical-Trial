"""
提取 PPTX 每页的正文/表格/备注/元素构成，输出为 JSON，供撰写演讲逐字稿使用。

用法:
    python extract_pptx_content_for_script.py <pptx路径> <输出json路径>
"""

import json
import sys
from pathlib import Path

from pptx import Presentation


def iter_shape_text(shape):
    """递归收集形状中的文本（含组合形状）。"""
    texts = []

    if shape.shape_type == 6:  # GROUP
        for s in shape.shapes:
            texts.extend(iter_shape_text(s))
        return texts

    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.replace("\n", " ").strip() for c in row.cells]
            texts.append(" | ".join(cells))
        return texts

    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            texts.append(t)
    return texts


def slide_inventory(slide):
    has_table = False
    has_picture = False
    has_chart = False
    texts = []
    for shape in slide.shapes:
        if getattr(shape, "has_table", False) and shape.has_table:
            has_table = True
        if shape.shape_type == 13:  # PICTURE
            has_picture = True
        if getattr(shape, "has_chart", False) and shape.has_chart:
            has_chart = True
        texts.extend(iter_shape_text(shape))
    return texts, has_table, has_picture, has_chart


def main():
    pptx_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])

    prs = Presentation(str(pptx_path))
    data = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts, has_table, has_pic, has_chart = slide_inventory(slide)

        title = ""
        if slide.shapes.title is not None and slide.shapes.title.text.strip():
            title = slide.shapes.title.text.strip().replace("\n", " ")

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        body = [t for t in texts if t != title]
        data.append(
            {
                "page": idx,
                "title": title,
                "body": body,
                "has_table": has_table,
                "has_picture": has_pic,
                "has_chart": has_chart,
                "existing_notes": notes,
                "char_count": sum(len(t) for t in body),
            }
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    total = sum(d["char_count"] for d in data)
    print(f"slides={len(data)} total_body_chars={total} -> {out_path}")


if __name__ == "__main__":
    main()
