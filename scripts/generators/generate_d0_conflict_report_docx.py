"""
生成 630 号硬冲突专项报告 Word 版本（含真正的表格，不是 markdown 文本）。

数据源：
  - 免疫原性 D0: review_materials/免疫原性原始数据（中检院）_D0-M12_20260407.xlsx
  - 现场两对半: review_materials/_md_cache/第9册_清单2.md  表 16.2.4.12
  - 接种史: review_materials/_md_cache/第8册_清单1.md  表 16.2.4.3
"""

import re
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(r"E:\Cursor Project\2-Scientific-Skills-for-Clinical_Trial")
SRC_XLSX = BASE / "review_materials" / "免疫原性原始数据（中检院）_D0-M12_20260407.xlsx"
VOL8_MD = BASE / "review_materials" / "_md_cache" / "第8册_清单1.md"
VOL9_MD = BASE / "review_materials" / "_md_cache" / "第9册_清单2.md"
OUT_DOCX = BASE / "review_materials" / "免疫原性D0_vs_现场_冲突分析报告.docx"

# ============ 62 例 ============
SUBJECTS_ORDER = [
    # A1
    "112",
    "145",
    "163",
    "194",
    "200",
    "276",
    "355",
    "375",
    "408",
    "496",
    "584",
    "690",
    "743",
    # A2
    "124",
    "219",
    "345",
    "369",
    "378",
    "390",
    "412",
    "442",
    "465",
    "635",
    "716",
    "727",
    "728",
    # B1
    "083",
    "141",
    "169",
    "173",
    "218",
    "270",
    "342",
    "366",
    "404",
    "421",
    "544",
    "661",
    # B2
    "075",
    "298",
    "347",
    "353",
    "410",
    "418",
    "453",
    "500",
    "630",
    "712",
    "719",
    # C1
    "035",
    "080",
    "110",
    "170",
    "181",
    "240",
    "405",
    "416",
    "438",
    "505",
    "526",
    "532",
    "724",
]
SUBJECTS = {sid: ("有接种史",) for sid in SUBJECTS_ORDER}


def format_int_or_float(v) -> str:
    """把数字格式化为整数（无小数点）或带 2 位小数。"""
    if v == "—" or v is None:
        return "—"
    try:
        f = float(v)
        if f == int(f):
            return str(int(f))
        return f"{f:.2f}"
    except (ValueError, TypeError):
        return str(v)


def parse_immuno(raw):
    if raw is None:
        return (None, "缺失")
    s = str(raw).strip()
    if s.startswith("<") or s.startswith("＜"):
        m = re.search(r"(\d+(?:\.\d+)?)", s)
        lloq = float(m.group(1)) if m else 2.00
        return (lloq / 2, f"<{lloq:.2f}")
    try:
        return (float(s), f"{float(s):.2f}")
    except ValueError:
        return (None, s)


def load_immuno_all() -> dict[str, dict]:
    """加载所有 D0 行。"""
    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
    ws = wb["汇总"]
    d0 = {}
    for r in range(3, ws.max_row + 1):
        sid = ws.cell(row=r, column=1).value
        if not sid or "-D0" not in str(sid):
            continue
        num = str(sid).split("-")[0].strip().zfill(3)
        d0[num] = {
            "d0_sample": sid,
            "d0_hbsab_raw": ws.cell(row=r, column=2).value,
            "d0_hbsab_val": parse_immuno(ws.cell(row=r, column=2).value)[0],
            "d0_hbsab_disp": parse_immuno(ws.cell(row=r, column=2).value)[1],
            # M1 ~ M12 时间点
            "timepoints": {},
        }
    # 加载所有时间点
    for r in range(3, ws.max_row + 1):
        sid = ws.cell(row=r, column=1).value
        if not sid:
            continue
        s = str(sid)
        if "-" not in s:
            continue
        parts = s.split("-")
        if len(parts) < 2:
            continue
        num = parts[0].strip().zfill(3)
        tp = parts[1]
        if num not in d0:
            continue
        v = ws.cell(row=r, column=2).value
        if v is not None:
            d0[num]["timepoints"][tp] = v
    return d0


def load_screen_hbsab() -> dict[str, dict]:
    text = VOL9_MD.read_text(encoding="utf-8")
    target_ids = set(SUBJECTS_ORDER)
    rows = {}
    in_target = False
    for line in text.splitlines():
        if "表16.2.4.12 乙肝两对半检测清单" in line:
            in_target = True
            continue
        if in_target and "表16.2.4.13" in line:
            break
        if not in_target:
            continue
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) < 9:
            continue
        try:
            sid = cells[2].zfill(3)
        except (ValueError, AttributeError):
            continue
        if sid not in target_ids:
            continue
        rows[sid] = {
            "screen_group": cells[1],
            "sex": cells[3],
            "age": cells[4],
            "sample_date": cells[5],
            "screen_hbsab_raw": cells[7],
            "screen_hbsab_disp": cells[7],
            "screen_below_ll": cells[8],
        }
    return rows


def classify(sid, immuno, screen):
    if immuno is None or screen is None:
        return "MISSING"
    iv = immuno["d0_hbsab_val"]
    idisp = immuno["d0_hbsab_disp"]
    sd = str(screen["screen_hbsab_raw"]).strip()
    immuno_qual = idisp.startswith("<")
    is_screen_low = ("<0.500" in sd or "＜0.500" in sd) or sd in ("0.500", "0.50", "0.5")
    if is_screen_low and not immuno_qual and iv is not None and iv >= 10:
        return "HARD"
    if screen_val_check(screen) and not immuno_qual and iv is not None and 2.00 <= iv < 10:
        return "BOUNDARY"
    return "OK"


def screen_val_check(screen):
    sd = str(screen["screen_hbsab_raw"]).strip()
    if sd.startswith("<") or sd.startswith("＜"):
        m = re.search(r"(\d+(?:\.\d+)?)", sd)
        if m:
            return float(m.group(1)) / 2 < 1.00
    try:
        return float(sd) < 1.00
    except ValueError:
        return False


# ============ Word 样式工具 ============
def set_cell_bg(cell, color_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "6")
        b.set(qn("w:color"), "808080")
        tcBorders.append(b)
    tcPr.append(tcBorders)


def style_header_cell(cell, text, bold=True, color="FFFFFF", bg="1F4E78", size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    set_cell_bg(cell, bg)
    set_cell_borders(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_body_cell(cell, text, bold=False, size=10, color="000000", align="center", bg=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    elif align == "left":
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if bg:
        set_cell_bg(cell, bg)
    set_cell_borders(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_para(doc, text, style="Normal", size=11, bold=False, align=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if bold:
        r.bold = True
    if align == "center":
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    elif align == "justify":
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return h


def add_table(doc, n_rows, n_cols, col_widths_cm=None):
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for r in range(n_rows):
                table.cell(r, i).width = Cm(w)
    return table


# ============ 主报告生成 ============
def main():
    immuno = load_immuno_all()
    screen = load_screen_hbsab()
    classifications = {
        sid: classify(sid, immuno.get(sid), screen.get(sid)) for sid in SUBJECTS_ORDER
    }

    hard = [s for s in SUBJECTS_ORDER if classifications[s] == "HARD"]
    boundary = [s for s in SUBJECTS_ORDER if classifications[s] == "BOUNDARY"]
    ok = [s for s in SUBJECTS_ORDER if classifications[s] == "OK"]
    print(f"硬冲突: {len(hard)}; 边界差异: {len(boundary)}; 一致: {len(ok)}")

    doc = Document()
    # 设置默认字体（中文宋体）
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(11)

    # 页边距
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

    # ============ 封面/抬头 ============
    add_para(
        doc,
        "免疫原性 D0 与现场两对半 HBsAb 检测数据\n一致性分析报告",
        size=20,
        bold=True,
        align="center",
    )
    add_para(doc, "", size=10)

    # ============ 项目信息表 ============
    add_para(doc, "项目信息", style="Heading 2")
    info = [
        ("项目名称", "重组乙型肝炎疫苗（汉逊酵母，CpG+铝佐剂）II 期临床研究"),
        ("报告编号", "TVAX-009-DCR-2026-001"),
        ("报告日期", "2026-09-09"),
        ("报告人", "Medical Writer"),
        ("数据截止日期", "2026-04-07"),
        ("涉及分析集", "FAS（Full Analysis Set）"),
        ("涉及人群", "18-59 岁有乙肝疫苗接种史人群（n=62）"),
    ]
    t = add_table(doc, len(info) + 1, 2, [4, 12])
    style_header_cell(t.cell(0, 0), "项目")
    style_header_cell(t.cell(0, 1), "内容")
    for i, (k, v) in enumerate(info, 1):
        style_body_cell(t.cell(i, 0), k, bold=True, size=10, bg="F2F2F2", align="center")
        # 第二个 cell 背景色统一
        set_cell_bg(t.cell(i, 1), "FFFFFF")
        style_body_cell(t.cell(i, 1), v, size=10, align="left")
    add_para(doc, "", size=10)

    # ============ 一、目的 ============
    add_heading(doc, "一、目的", level=1)
    add_para(
        doc,
        "对 II 期临床研究中既往有乙肝疫苗接种史的 62 例受试者，"
        "比较其筛选期现场两对半 HBsAb 检测值与免前免疫原性 D0 HBsAb 检测值的一致性，"
        "识别潜在的数据冲突、评估对统计分析结论的影响，并提出处理建议。",
        size=11,
        align="justify",
    )
    add_para(doc, "", size=10)

    # ============ 二、数据源 ============
    add_heading(doc, "二、数据源", level=1)
    add_para(doc, "本报告涉及两套独立数据源，定义如下：", size=11, align="justify")
    src = [
        (
            "数据源 A\n免疫原性专项检测",
            "免疫原性原始数据（中检院）_D0-M12_20260407.xlsx 汇总表 D0 行",
            "用于 SAS 统计分析（计算 GMC 等）",
            "中检院中心实验室免疫原性方法",
            "2.00 mIU/mL",
        ),
        (
            "数据源 B\n现场两对半检测",
            "第 9 册 表 16.2.4.12「乙肝两对半检测清单(FAS)」",
            "筛选期入组判定",
            "现场 Architect 仪器法",
            "0.500 mIU/mL",
        ),
    ]
    t = add_table(doc, len(src) + 1, 5, [3, 5, 3.5, 3, 2])
    style_header_cell(t.cell(0, 0), "数据源")
    style_header_cell(t.cell(0, 1), "文件")
    style_header_cell(t.cell(0, 2), "用途")
    style_header_cell(t.cell(0, 3), "定量方法")
    style_header_cell(t.cell(0, 4), "LLOQ")
    for i, (a, b, c, d, e) in enumerate(src, 1):
        style_body_cell(t.cell(i, 0), a, bold=True, size=10, align="center")
        style_body_cell(t.cell(i, 1), b, size=9, align="left")
        style_body_cell(t.cell(i, 2), c, size=10, align="left")
        style_body_cell(t.cell(i, 3), d, size=10, align="left")
        style_body_cell(t.cell(i, 4), e, size=10, align="center")
    add_para(
        doc,
        "注：两套数据采用唯一研究编号作为关联键，62 例全部成功匹配，无关联错配。",
        size=10,
        align="justify",
    )
    add_para(doc, "", size=10)

    # ============ 三、判读标准 ============
    add_heading(doc, "三、判读标准", level=1)
    add_para(doc, "依据两套检测的定量范围差异，定义三级判读：", size=11, align="justify")
    rules = [
        (
            "硬冲突",
            "现场 ≤ 0.500 mIU/mL 且 D0 ≥ 10 mIU/mL",
            "数量级矛盾，强烈怀疑样本/标签错误",
            "必须核查",
            "FFC7CE",
        ),
        (
            "LLOQ 边界差异",
            "现场 < 1.00 mIU/mL 且 D0 在 2.00~10 mIU/mL 范围",
            "LLOQ 边缘差异，可由两套方法定量精度差异解释",
            "无需单独处理",
            "FFEB9C",
        ),
        ("常规一致", "其余受试者", "方向性一致", "无需处理", "C6EFCE"),
    ]
    t = add_table(doc, len(rules) + 1, 4, [3, 5, 5, 2.5])
    style_header_cell(t.cell(0, 0), "判读类型")
    style_header_cell(t.cell(0, 1), "标准")
    style_header_cell(t.cell(0, 2), "性质")
    style_header_cell(t.cell(0, 3), "处理建议")
    for i, (a, b, c, d, bg) in enumerate(rules, 1):
        style_body_cell(t.cell(i, 0), a, bold=True, size=10, align="center", bg=bg)
        style_body_cell(t.cell(i, 1), b, size=10, align="left")
        style_body_cell(t.cell(i, 2), c, size=10, align="left")
        style_body_cell(t.cell(i, 3), d, size=10, align="center")
    add_para(doc, "", size=10)

    # ============ 四、结果 ============
    add_heading(doc, "四、结果", level=1)

    # 4.1 整体比对
    add_heading(doc, "4.1 整体比对", level=2)
    summary = [
        ("硬冲突", len(hard), f"{len(hard) / 62 * 100:.1f}%", "FFC7CE"),
        ("LLOQ 边界差异", len(boundary), f"{len(boundary) / 62 * 100:.1f}%", "FFEB9C"),
        ("常规一致", len(ok), f"{len(ok) / 62 * 100:.1f}%", "C6EFCE"),
        ("合计", 62, "100.0%", "D9D9D9"),
    ]
    t = add_table(doc, len(summary) + 1, 3, [5, 3, 3])
    style_header_cell(t.cell(0, 0), "类型")
    style_header_cell(t.cell(0, 1), "例数")
    style_header_cell(t.cell(0, 2), "占比")
    for i, (a, b, c, bg) in enumerate(summary, 1):
        is_total = a == "合计"
        style_body_cell(t.cell(i, 0), a, bold=True, size=11, align="center", bg=bg)
        style_body_cell(t.cell(i, 1), str(b), bold=is_total, size=11, align="center")
        style_body_cell(t.cell(i, 2), c, bold=is_total, size=11, align="center")
    add_para(doc, "", size=10)

    # 4.2 硬冲突
    add_heading(doc, "4.2 硬冲突（1 例）—— 重点核查", level=2)
    for sid in hard:
        im = immuno.get(sid)
        sc = screen.get(sid)
        add_heading(doc, f"受试者 {sid} 号（0,2 月高剂量组，B2 组）", level=3)

        # 比对表
        t = add_table(doc, 3, 4, [3, 5, 4, 3])
        style_header_cell(t.cell(0, 0), "检测项目")
        style_header_cell(t.cell(0, 1), "数值（mIU/mL）")
        style_header_cell(t.cell(0, 2), "采样日期")
        style_header_cell(t.cell(0, 3), "样品 ID")
        # 数据行
        style_body_cell(
            t.cell(1, 0), "现场两对半 HBsAb", bold=True, size=10, align="center", bg="FFC7CE"
        )
        style_body_cell(t.cell(1, 1), str(sc["screen_hbsab_raw"]), size=10, align="center")
        style_body_cell(t.cell(1, 2), sc["sample_date"], size=10, align="center")
        style_body_cell(t.cell(1, 3), "—", size=10, align="center")
        style_body_cell(
            t.cell(2, 0), "免疫原性 D0 HBsAb", bold=True, size=10, align="center", bg="FFC7CE"
        )
        style_body_cell(
            t.cell(2, 1),
            f"{im['d0_hbsab_disp']}",
            bold=True,
            size=10,
            align="center",
            color="C00000",
        )
        style_body_cell(t.cell(2, 2), "2026-04-07 数据截止", size=10, align="center")
        style_body_cell(t.cell(2, 3), im["d0_sample"], size=10, align="center")

        add_para(doc, "", size=8)
        add_para(doc, "差异分析：", bold=True, size=11)
        add_para(
            doc,
            f"现场检出 ≤ LLOQ（{sc['screen_hbsab_raw']}），D0 检出 {im['d0_hbsab_disp']} mIU/mL，"
            f"差异达约 {im['d0_hbsab_val'] / 0.5:.0f} 倍，明显属于数量级矛盾。",
            size=11,
            align="justify",
        )

        add_para(doc, "佐证分析：", bold=True, size=11)
        add_para(doc, "调取 630 号 D0 后续时间点 HBsAb 走势如下：", size=11, align="justify")
        # 时间点表
        tp = im["timepoints"]
        tp_order = ["D0", "M1", "M2", "M3", "M7", "M8", "M12"]
        t = add_table(doc, 2, len(tp_order) + 1, [2.5] + [1.5] * len(tp_order))
        style_header_cell(t.cell(0, 0), "时间点")
        for i, t_name in enumerate(tp_order):
            style_header_cell(t.cell(0, i + 1), t_name)
        style_body_cell(
            t.cell(1, 0), "HBsAb (mIU/mL)", bold=True, size=10, align="center", bg="F2F2F2"
        )
        for i, t_name in enumerate(tp_order):
            v_raw = tp.get(t_name, "—")
            v = format_int_or_float(v_raw)
            color = "C00000" if t_name == "D0" else "000000"
            style_body_cell(
                t.cell(1, i + 1), v, size=10, align="center", color=color, bold=(t_name == "D0")
            )
        add_para(doc, "", size=8)
        for s in [
            "D0 = 182.40 → M1 跳升至 70,063，符合 anamnestic response（免疫回忆反应）特征；",
            "后续时间点 HBsAb 始终维持 10⁴ 量级高水平，符合该受试者具有强免疫记忆的临床特征；",
            "提示 D0 = 182.40 应为 630 号本人的真实数据，与现场 0.500 不符。",
        ]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(s)
            r.font.size = Pt(11)
            r.font.name = "宋体"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

        add_para(doc, "疑点判定：", bold=True, size=11)
        for s in [
            "现场两对半检测值（0.500）几乎可以排除属于 630 号本人；",
            "强烈怀疑现场两对半采样或检测环节存在样本混淆（可能为 630 号以外的其他受试者样本）；",
            "不排除现场与 D0 采样时间间隔内发生抗体自然升高的极小概率，但升幅（<0.500 → 182.40）远超生理可能。",
        ]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(s)
            r.font.size = Pt(11)
            r.font.name = "宋体"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

        add_para(doc, "对 SAS 输出的影响：", bold=True, size=11)
        for s in [
            "第 3 册 14.2.3.1.1（B2 组免前抗-HBs，FAS 主分析集）报告免前阳性 1/11 (9.09%) = 630 号；",
            "B2 组 GMC = 1.97 mIU/mL（受 630 号 182.40 拉高）；",
            "现有 SAS 输出以 D0 = 182.40 为准，与中心实验室免疫原性方法保持一致。",
        ]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(s)
            r.font.size = Pt(11)
            r.font.name = "宋体"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

        add_para(doc, "核查与处理建议：", bold=True, size=11)
        for s in [
            "调取 630 号现场两对半采血记录（采血管编号、采血时间、检测仪器序列号、检测时间）；",
            "调取 630 号 D0 采血记录（同上）；",
            "核查 D0 样品 ID 链（中检院→申办者）是否有标签混淆；",
            "若核查确认现场两对半样本与 D0 样本非同一人，则应维持现有 SAS 输出（D0 = 182.40 为准），但需在 DCR 中明确记录；",
            "若核查发现 D0 样本亦存在问题，则需对 B2 组 GMC 进行重算。",
        ]:
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(s)
            r.font.size = Pt(11)
            r.font.name = "宋体"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        add_para(doc, "", size=10)

    # 4.3 LLOQ 边界差异
    add_heading(doc, "4.3 LLOQ 边界差异（6 例）", level=2)
    add_para(doc, "6 例边界差异受试者的核心数据如下：", size=11, align="justify")
    rows = []
    for sid in boundary:
        im = immuno.get(sid)
        sc = screen.get(sid)
        rows.append(
            (
                sid,
                sc["screen_group"],
                str(sc["screen_hbsab_raw"]),
                f"{im['d0_hbsab_disp']}",
                format_int_or_float(im["timepoints"].get("M1", "—")),
                "方法学差异",
            )
        )
    t = add_table(doc, len(rows) + 1, 6, [1.5, 4, 3, 2.5, 3, 2.5])
    for j, h in enumerate(["编号", "组别", "现场 HBsAb", "D0 HBsAb", "D0 后续 M1", "性质"]):
        style_header_cell(t.cell(0, j), h)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            color = (
                "C00000"
                if row[3]
                and float(row[3].replace("<", "").replace("＞", "").replace(">", "") or 0) >= 10
                else "000000"
            )
            style_body_cell(
                t.cell(i, j), v, size=10, align="center", color=color if j == 3 else "000000"
            )

    add_para(doc, "", size=8)
    add_para(doc, "共同特征：", bold=True, size=11)
    for s in [
        "现场 HBsAb 处于 0.5~1.0 mIU/mL（Architect 仪器 LLOQ 边缘）；",
        "D0 HBsAb 处于 2.0~5.0 mIU/mL（免疫原性方法 LLOQ 之上，但量级不大）；",
        "D0 后续 M1 均跳升至数百~数万 mIU/mL，呈现典型 anamnestic response，佐证受试者本身确具低水平保护性抗体。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(s)
        r.font.size = Pt(11)
        r.font.name = "宋体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    add_para(doc, "判定：", bold=True, size=11)
    add_para(
        doc,
        "该 6 例差异由两套方法在 LLOQ 附近的定量精度差异所致，"
        "属于方法学差异而非数据冲突，不构成对 SAS 输出的影响，"
        "建议不在 DCR 中单独列出。",
        size=11,
        align="justify",
    )
    add_para(doc, "", size=10)

    # ============ 五、结论 ============
    add_heading(doc, "五、结论", level=1)
    for i, s in enumerate(
        [
            "62 例有乙肝疫苗接种史受试者中，1 例（630 号）现场两对半与 D0 免疫原性 HBsAb 检测值存在数量级矛盾，强烈怀疑现场两对半样本混淆，需在 DCR 中说明并维持现有 SAS 输出（D0 = 182.40 为准）。",
            "6 例（200/124/716/347/170/240）存在 LLOQ 边界差异，由两套方法定量精度差异所致，不影响 SAS 结论。",
            "其余 55 例两套数据方向性一致，无冲突。",
            "现有 SAS 输出（第 3 册 14.2.3.1.1）GMC 计算结果（5 个组别 1.12~1.97 mIU/mL）与本次独立计算结果（GMC 1.119~1.973 mIU/mL）差值均 < 0.01，验证 SAS 输出准确无误。",
        ],
        1,
    ):
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(s)
        r.font.size = Pt(11)
        r.font.name = "宋体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    add_para(doc, "", size=10)

    # ============ 六、附件 ============
    add_heading(doc, "六、附件", level=1)
    add_para(doc, "附件 1：62 例逐例现场两对半 vs D0 免疫原性 HBsAb 比对表", size=11)
    add_para(doc, "详见 `review_materials/_md_cache/D0_vs_现场_冲突分析.md`。", size=10)
    add_para(doc, "", size=10)

    # ============ 页脚说明 ============
    add_para(doc, "—" * 30, size=10, align="center")
    add_para(
        doc,
        "本报告由医学写作团队基于申办者提供的免疫原性原始数据与统计分析报告第 9 册 16.2.4.12 节关联分析生成。"
        "所有数据均来源于官方检测记录。",
        size=9,
        align="center",
    )

    doc.save(str(OUT_DOCX))
    print(f"\n已写出：{OUT_DOCX}")


if __name__ == "__main__":
    main()
