# -*- coding: utf-8 -*-
"""
自检脚本：Ⅱ期 老年≥60岁 两剂次桥接非劣效分析 交付文件复核
独立于生成脚本，重新走一遍数据源与算法，逐项比对待交付 Excel。

检查项：
  A. 工作簿结构（子表齐全、每个比较子表含 PPS 与 FAS 两行）
  B. 表内数值自洽（率 = n/N；率差 = 率1 − 率2；判定 = CI下限 > −5%）
  C. 独立重算 PPS（按受试者 ID 去重后计数，与生成脚本逐行计数互为印证）
  D. 独立重读 TFL 官方 FAS 汇总表（第3册 表14.2.4.2.1），核对官方发表值
  E. 独立重读 TFL 官方 PPS 汇总表，核对 PPS 发表值
  F. 说明页文本完整性（各节标题齐全、无占位符）
"""

import os
import glob
import re

import docx
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TFL_DIR = os.path.join(BASE, "review_materials", "TFL-Phase2")
OUT_XLSX = os.path.join(
    BASE, "outputs",
    "TVAX-009-002(II)二期_老年≥60岁_两剂次桥接非劣效分析_PPS+FAS.xlsx",
)

C3 = "0,1,6月高剂量组(C3)"
C2 = "阳性对照组(C2)"
V2, V3, V7, V8 = ("首剂接种后2个月", "首剂接种后3个月",
                  "首剂接种后7个月", "首剂接种后8个月")
MARGIN = -5.0

FAILS = []


def chk(cond, msg):
    print(f"  [{'PASS' if cond else 'FAIL'}] {msg}")
    if not cond:
        FAILS.append(msg)
    return cond


# ---------------------------------------------------------------------------
# 独立重算 PPS：按受试者 ID 去重
# ---------------------------------------------------------------------------
DIAG = {}


def independent_pps_count():
    path = glob.glob(os.path.join(TFL_DIR, "*清单2*.docx"))[0]
    d = docx.Document(path)
    t = d.tables[9]
    hdr = [x.text.strip().replace("\n", "") for x in t.rows[0].cells]
    DIAG["header"] = hdr
    # 按 (受试者ID, 组别, 访视) 聚合，避免同一受试者重复行干扰
    agg = {}
    sample = []
    for ri, r in enumerate(t.rows[1:], start=1):
        c = [x.text.strip().replace("\n", "") for x in r.cells]
        try:
            age = int(c[4])
        except ValueError:
            continue
        if age < 60:
            continue
        if len(sample) < 3:
            sample.append(c[:11])
        key = (c[2], c[1], c[5])  # 研究编号 + 组别 + 访视
        ev = (c[9] == "是") or (c[10] == "是")
        if key not in agg:
            agg[key] = ev
        else:
            agg[key] = agg[key] or ev
    DIAG["n_rows_ge60"] = None
    DIAG["n_uniq_keys"] = len(agg)
    DIAG["sample"] = sample
    DIAG["uniq_sid"] = len({k[0] for k in agg})
    out = {}
    for (sid, grp, vis), ev in agg.items():
        k = (grp, vis)
        x, n = out.get(k, (0, 0))
        out[k] = (x + (1 if ev else 0), n + 1)  # (阳转数, 合计)
    return out


# ---------------------------------------------------------------------------
# 在原文档中定位标题行所属表格
# ---------------------------------------------------------------------------
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def find_table_by_caption(d, keyword):
    """按题注文字定位其后出现的第一个表格，返回 (tables 下标, 题注全文)。

    注意：python-docx 的 Table 对象不能用 list.index 比较，只能按文档流顺序计数。
    """
    ti = -1
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "tbl":
            ti += 1
        elif tag == "p":
            txt = "".join(n.text or "" for n in el.iter(W_NS + "t"))
            if keyword in txt:
                # 题注位于表格之前 → 该表格的序号为当前计数 +1
                return ti + 1, txt.strip()
    return None, None


def read_official_table(path, caption_kw):
    d = docx.Document(path)
    idx, cap = find_table_by_caption(d, caption_kw)
    if idx is None:
        return None, None
    t = d.tables[idx]
    rows = []
    for r in t.rows:
        rows.append([x.text.strip().replace("\n", "") for x in r.cells])
    return rows, cap


MONTH2VIS = {i: f"首剂接种后{i}个月" for i in range(1, 13)}


def parse_official(rows):
    """解析官方汇总表（第3册 表14.2.4.2.x 系列）。

    表结构：
      第 0 行表头：['', '<组别1>(N=xx)', '<组别2>(N=xx)', '率差 (95% CI)(%)[2]', 'P值[2]']
      其后按时间点分块：
        块标题行    col0 = 【首剂免后N个月抗-HBs阳转率】
        阳转行      col0 = 阳转 n (%)，col1/col2 = 【n (率)】
        CI 行       col0 = 95%可信区间 (%)[1]，col1/col2 = 【下限, 上限】
        合计行      col0 = 合计 (Missing)
      跨时点比较行：col0 = 【C3组首剂免后X个月 VS C2组首剂免后Y个月】，col3 = 率差 (下限, 上限)

    返回 (单元字典 {(组别, 访视): (x, N)}, 跨时点字典 {(月X, 月Y): (率差, 下限, 上限)})
    """
    cells, cross = {}, {}
    if not rows:
        return cells, cross

    def grp_of(txt):
        for g in (C3, C2):
            if txt.startswith(g):
                return g
        return None

    header = rows[0]
    g1 = grp_of(header[1]) if len(header) > 1 else None
    g2 = grp_of(header[2]) if len(header) > 2 else None
    cur_month = None
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        c0 = r[0]
        m = re.search(r"首剂免后(\d+)个月", c0)
        if "VS" in c0:
            # 跨时点比较行：C3组首剂免后X个月 VS C2组首剂免后Y个月
            mm = re.search(r"首剂免后(\d+)个月\s*VS\s*C2组首剂免后(\d+)个月", c0)
            if mm and len(r) > 3:
                dd = re.search(r"(-?[\d.]+)\s*[\(（]\s*(-?[\d.]+)\s*,\s*(-?[\d.]+)\s*[\)）]", r[3])
                if dd:
                    cross[(int(mm.group(1)), int(mm.group(2)))] = (
                        float(dd.group(1)), float(dd.group(2)), float(dd.group(3)))
            continue
        if c0.startswith("阳转") and cur_month:
            for g, col in ((g1, 1), (g2, 2)):
                if not g or col >= len(r):
                    continue
                mm = re.match(r"\s*(\d+)\s*[\(（]([\d.]+)[\)）]", r[col])
                if mm:
                    cells[(g, MONTH2VIS[cur_month])] = (int(mm.group(1)), None)
        elif c0.startswith("合计") and cur_month:
            for g, col in ((g1, 1), (g2, 2)):
                if not g or col >= len(r):
                    continue
                mm = re.match(r"\s*(\d+)", r[col])
                if mm:
                    x, _ = cells.get((g, MONTH2VIS[cur_month]), (None, None))
                    cells[(g, MONTH2VIS[cur_month])] = (x, int(mm.group(1)))
        elif m:
            cur_month = int(m.group(1))
    return cells, cross


# ---------------------------------------------------------------------------
def main():
    print("=" * 78)
    print("交付文件自检")
    print("=" * 78)
    print(f"文件：{OUT_XLSX}")
    print()

    wb = load_workbook(OUT_XLSX, data_only=True)
    print("[A] 工作簿结构")
    sheets = wb.sheetnames
    print(f"     子表：{sheets}")
    chk("00_说明" in sheets, "存在说明页 00_说明")
    chk("01_汇总" in sheets, "存在汇总页 01_汇总")
    comp_sheets = [s for s in sheets if s.startswith("比较")]
    chk(len(comp_sheets) == 3, f"存在 3 个比较子表（实际 {len(comp_sheets)}）")
    print()

    print("[B] 表内数值自洽与判定一致性")
    expected = {
        "比较1_试验M3vs对照M7": (V3, V7),
        "比较2_试验M2vs对照M7": (V2, V7),
        "比较3_试验M3vs对照M8": (V3, V8),
    }
    table_vals = {}
    for s in comp_sheets:
        ws = wb[s]
        got = []
        for rr in (9, 10):
            a = ws.cell(row=rr, column=1).value
            if not a:
                continue
            n1 = ws.cell(row=rr, column=2).value
            p1 = ws.cell(row=rr, column=3).value
            n2 = ws.cell(row=rr, column=6).value
            p2 = ws.cell(row=rr, column=7).value
            dif = ws.cell(row=rr, column=10).value
            dlo = ws.cell(row=rr, column=11).value
            dhi = ws.cell(row=rr, column=12).value
            vd = ws.cell(row=rr, column=13).value
            x1, N1 = [int(z) for z in n1.split("/")]
            x2, N2 = [int(z) for z in n2.split("/")]
            ok_rate = abs(p1 - round(x1 / N1 * 100, 2)) < 0.011 and \
                      abs(p2 - round(x2 / N2 * 100, 2)) < 0.011
            ok_dif = abs(dif - round(p1 - p2, 2)) < 0.011
            ok_vd = (vd.startswith("成立") == (dlo > MARGIN))
            ok_ord = dlo < dif < dhi
            chk(ok_rate, f"{s} / {a} 率 = n/N 自洽（{n1}→{p1}%，{n2}→{p2}%）")
            chk(ok_dif, f"{s} / {a} 率差 = 率1 − 率2（{dif}%）")
            chk(ok_vd, f"{s} / {a} 判定与 CI 下限一致（下限 {dlo}，判定 {vd}）")
            chk(ok_ord, f"{s} / {a} CI 区间有序（{dlo} < {dif} < {dhi}）")
            got.append((a, x1, N1, x2, N2))
        sets = [g[0] for g in got]
        chk(sets == ["PPS", "FAS"], f"{s} 同时含 PPS 与 FAS 两行（实际 {sets}）")
        table_vals[s] = got
    print()

    print("[C] 独立重算 PPS（按受试者 ID 去重计数）")
    pps = independent_pps_count()
    print(f"     表头：{DIAG['header']}")
    print(f"     去重后唯一键 {DIAG['n_uniq_keys']} 个，唯一受试者 ID {DIAG['uniq_sid']} 个")
    for s in DIAG["sample"]:
        print(f"     样例行：{s}")
    print(f"     解析结果：{pps}")
    for s in comp_sheets:
        vt, vc = expected[s]
        x1, N1 = pps.get((C3, vt), (None, None))
        x2, N2 = pps.get((C2, vc), (None, None))
        got = table_vals[s]
        m = {g[0]: (g[1], g[2], g[3], g[4]) for g in got}
        chk(m.get("PPS") == (x1, N1, x2, N2),
            f"{s} PPS 与独立重算一致（表内 {m.get('PPS')} vs 重算 {(x1, N1, x2, N2)}）")
    print()

    print("[D/E] 独立重读 TFL 官方汇总表并与交付文件对账")
    p3 = glob.glob(os.path.join(TFL_DIR, "*免疫原性2*.docx"))[0]
    print(f"     文件：{os.path.basename(p3)}")

    SET_OF = {"14.2.4.2.1": "FAS", "14.2.4.2.3": "PPS", "14.2.4.2.4": "PPS",
              "14.2.4.2.6": "PPS", "14.2.4.2.7": "PPS"}
    official = {}
    cross_all = {}
    for kw in ("14.2.4.2.1", "14.2.4.2.3", "14.2.4.2.4", "14.2.4.2.6", "14.2.4.2.7"):
        rows, cap = read_official_table(p3, kw)
        if rows is None:
            print(f"     （未定位到表 {kw}）")
            continue
        cells, cross = parse_official(rows)
        print(f"     {kw} [{SET_OF[kw]}]  单元 {len(cells)} 个，跨时点 {len(cross)} 个")
        for k in sorted(cells):
            official.setdefault(k, set()).add(cells[k])
        for k, v in cross.items():
            cross_all[(SET_OF[kw],) + k] = v

    # 交付文件中记录的数值（PPS 来自表内，FAS 来自官方值）
    delivered = {}
    for s in comp_sheets:
        for a, x1, N1, x2, N2 in table_vals[s]:
            delivered[(a, C3, expected[s][0])] = (x1, N1)
            delivered[(a, C2, expected[s][1])] = (x2, N2)

    print()
    print("     —— 单组 n/N 对账（交付文件 vs 独立重算 vs 官方汇总表）——")
    for (a, g, vis), v in sorted(delivered.items()):
        rec = pps.get((g, vis)) if a == "PPS" else None
        off = official.get((g, vis))
        ok_rec = (rec is None) or (rec == v)
        ok_off = (off is None) or (v in off)
        note = []
        note.append(f"重算 {rec}" if rec else "重算 —")
        note.append(f"官方 {sorted(off)}" if off else "官方 —")
        chk(ok_rec and ok_off, f"{a} {g[:14]} {vis}：交付 {v}｜{'｜'.join(note)}")

    print()
    print("     —— 官方已发表的跨时点率差 95%CI 对账（交付文件计算值 vs 官方值）——")
    if not cross_all:
        print("     （官方汇总表中未检索到跨时点比较行）")
    for (aset, mo1, mo2), off in sorted(cross_all.items()):
        mine = None
        for s in comp_sheets:
            if expected[s] != (MONTH2VIS[mo1], MONTH2VIS[mo2]):
                continue
            for a, x1, N1, x2, N2 in table_vals[s]:
                if a != aset:
                    continue
                ws = wb[s]
                rr = 9 if a == "PPS" else 10
                mine = (ws.cell(row=rr, column=10).value,
                        ws.cell(row=rr, column=11).value,
                        ws.cell(row=rr, column=12).value)
        print(f"     官方 {aset}【C3 M{mo1} vs C2 M{mo2}】= {off}")
        if mine is None:
            continue
        ok = all(abs(float(mine[i]) - off[i]) < 0.02 for i in range(3))
        chk(ok, f"{aset} C3 M{mo1} vs C2 M{mo2} 率差(下限, 上限) 与官方一致"
                f"（交付 {mine} vs 官方 {off}）")
    print()

    print("[F] 说明页文本完整性")
    ws = wb["00_说明"]
    texts = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v:
            texts.append(str(v))
    joined = "\n".join(texts)
    need = ["【核心结论速览】", "【一、分析目的】", "【二、数据来源】", "【三、人群与组别】",
            "【四、时点口径】", "【五、统计方法与非劣效判定】", "【六、各比较结果】",
            "【七、最终结论与依据】", "【八、局限性说明】"]
    for k in need:
        chk(k in joined, f"说明页含 {k}")
    chk("TODO" not in joined and "待填" not in joined, "说明页无占位符")
    chk(len(texts) >= 50, f"说明页正文行数充足（{len(texts)} 行）")
    print()

    print("=" * 78)
    if FAILS:
        print(f"自检未通过，共 {len(FAILS)} 项：")
        for f in FAILS:
            print("   - " + f)
    else:
        print("自检全部通过")
    print("=" * 78)


if __name__ == "__main__":
    main()
