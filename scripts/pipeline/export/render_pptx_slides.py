"""Render selected slides of a PPTX to PNG via PowerPoint COM.

Constraint (workspace policy): all generated .py must live under ``scripts/``.

Usage:
    python scripts/render_pptx_slides.py input.pptx 8 9 --outdir reports/009_render
"""

from __future__ import annotations

import argparse
from pathlib import Path


def render(pptx: Path, slides: list[int], outdir: Path, scale_w: int = 2000) -> list[Path]:
    import win32com.client as win32

    outdir = outdir.resolve()
    outdir.mkdir(parents=True, exist_ok=True)
    app = win32.gencache.EnsureDispatch("PowerPoint.Application")
    out: list[Path] = []
    try:
        pres = app.Presentations.Open(str(pptx), ReadOnly=True, Untitled=False, WithWindow=False)
        try:
            h = int(scale_w * pres.PageSetup.SlideHeight / pres.PageSetup.SlideWidth)
            for idx in slides:
                dst = outdir / f"slide{idx:02d}.png"
                pres.Slides.Item(idx).Export(str(dst), "PNG", scale_w, h)
                out.append(dst)
                print(f"OK {dst}")
        finally:
            pres.Close()
    finally:
        app.Quit()
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("slides", nargs="*", type=int, help="1-based slide numbers; omit for all")
    ap.add_argument("--outdir", default="reports/render")
    ap.add_argument("--width", type=int, default=2000)
    args = ap.parse_args()

    src = Path(args.pptx).resolve()
    slides = args.slides
    if not slides:
        import pptx as _p

        slides = list(range(1, len(_p.Presentation(str(src)).slides) + 1))
    render(src, slides, Path(args.outdir), args.width)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
