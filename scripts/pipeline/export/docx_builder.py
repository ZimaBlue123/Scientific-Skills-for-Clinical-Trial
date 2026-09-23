"""DOCX export utilities."""

from __future__ import annotations

import logging
from collections.abc import Iterable
from pathlib import Path

logger = logging.getLogger(__name__)

_ASCII_FONT = "Times New Roman"
_EAST_ASIA_FONT = "宋体"
_TARGET_STYLES = (
    "Normal",
    "Title",
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Table Grid",
)

def apply_cn_en_fonts(doc, styles: Iterable[str] | None = None) -> int:
    """Enforce document-wide fonts: Times New Roman for English, 宋体 for Chinese."""
    from docx.oxml.ns import qn

    target_styles = list(styles) if styles is not None else _TARGET_STYLES
    updated = 0

    styles_dict = getattr(doc, "styles", None)
    if styles_dict is None:
        return 0

    for style_name in target_styles:
        if style_name not in styles_dict:
            continue
        try:
            style = styles_dict[style_name]
            font = style.font
            font.name = _ASCII_FONT

            rpr = style.element.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            rfonts.set(qn("w:ascii"), _ASCII_FONT)
            rfonts.set(qn("w:hAnsi"), _ASCII_FONT)
            rfonts.set(qn("w:eastAsia"), _EAST_ASIA_FONT)
            rfonts.set(qn("w:cs"), _ASCII_FONT)
            updated += 1
        except (KeyError, ValueError):
            pass

    return updated

def create_clinical_docx(output_path: Path | str, content: str) -> None:
    """Create a basic clinical DOCX report with standard fonts."""
    from docx import Document

    doc = Document()
    apply_cn_en_fonts(doc)

    for line in content.split("\\n"):
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        else:
            doc.add_paragraph(line)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    logger.info(f"Created DOCX at {out}")

__all__ = ["create_clinical_docx", "apply_cn_en_fonts"]
