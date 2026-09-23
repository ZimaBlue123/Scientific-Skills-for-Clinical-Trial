"""PPTX export utilities."""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

def create_slide(prs, title: str, content: str, layout_idx: int = 1) -> None:
    """Create a slide with title and content."""
    layout = prs.slide_layouts[layout_idx]
    slide = prs.slides.add_slide(layout)

    if slide.shapes.title:
        slide.shapes.title.text = title

    for shape in slide.placeholders:
        if shape.placeholder_format.idx == 1:
            shape.text = content
            break

def inject_pptx_data(output_path: Path | str, data: dict[str, str]) -> None:
    """Generate a PPTX file from a dictionary of slide titles and contents."""
    from pptx import Presentation

    prs = Presentation()

    for title, content in data.items():
        create_slide(prs, title, content)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    logger.info(f"Created PPTX at {out}")

__all__ = ["inject_pptx_data", "create_slide"]
