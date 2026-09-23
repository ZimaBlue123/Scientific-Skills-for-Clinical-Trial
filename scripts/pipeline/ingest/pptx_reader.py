"""PPTX extraction module."""

from __future__ import annotations
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

def extract_pptx_text(pptx_path: Path | str) -> str:
    """Extract all text, including tables and group shapes, from a PPTX file."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    
    prs = Presentation(str(pptx_path))
    lines = []
    
    def _iter_shape_text(shape, depth=0):
        indent = "  " * depth
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines.append(f"{indent}[GROUP]")
            for sub in shape.shapes:
                _iter_shape_text(sub, depth + 1)
        elif shape.has_table:
            lines.append(f"{indent}[TABLE]")
            for row in shape.table.rows:
                cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                lines.append(f"{indent}  {' | '.join(cells)}")
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                lines.append(f"{indent}[TEXT] {text}")
                
    for i, slide in enumerate(prs.slides, start=1):
        lines.append(f"===== Slide {i} =====")
        for shape in slide.shapes:
            _iter_shape_text(shape)
        lines.append("")
        
    return "\n".join(lines)

def extract_pptx_notes(pptx_path: Path | str) -> list[str]:
    """Extract notes from all slides in a PPTX file."""
    from pptx import Presentation
    
    prs = Presentation(str(pptx_path))
    notes = []
    
    for slide in prs.slides:
        if slide.has_notes_slide:
            note_text = slide.notes_slide.notes_text_frame.text.strip()
            notes.append(note_text)
        else:
            notes.append("")
            
    return notes

__all__ = ["extract_pptx_text", "extract_pptx_notes"]
