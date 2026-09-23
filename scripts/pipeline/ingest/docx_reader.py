"""DOCX extraction module."""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

def extract_docx_text(docx_path: Path | str) -> str:
    """Extract all text from a .docx file."""
    from docx import Document

    doc = Document(str(docx_path))
    full_text: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))

    return "\\n\\n".join(full_text)

def extract_docx_tables(docx_path: Path | str) -> list[list[list[str]]]:
    """Extract structured tables from a .docx file."""
    from docx import Document

    doc = Document(str(docx_path))
    tables_data = []

    for table in doc.tables:
        table_data = []
        for row in table.rows:
            table_data.append([cell.text.strip() for cell in row.cells])
        tables_data.append(table_data)

    return tables_data

__all__ = ["extract_docx_text", "extract_docx_tables"]
