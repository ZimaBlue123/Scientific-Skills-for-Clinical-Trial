#!/usr/bin/env python3
"""Verify internal cross-references (chapter/section numbers) of a protocol .docx.

Extracts all Heading paragraphs and every in-text cross-reference of the form
"N.N[.N] 章节" / "详见 N.N" / "表 N-N", then reports references that do not match
any heading number found in the document.

Usage
-----
    python scripts/audit_protocol_crossrefs.py <protocol.docx>
"""

from __future__ import annotations

import re
import sys
from pathlib import Path


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: audit_protocol_crossrefs.py <protocol.docx>")
        return 2
    from docx import Document

    path = Path(sys.argv[1])
    doc = Document(str(path))

    headings: list[tuple[str, str]] = []
    for p in doc.paragraphs:
        style = (p.style.name or "") if p.style is not None else ""
        text = p.text.strip()
        if not text:
            continue
        is_heading_style = style.lower().startswith("heading") or style.lower().startswith("标题")
        # fall back: many Chinese protocols use custom styles, detect by numbering prefix
        is_numbered = bool(re.match(r"^\d+(?:\.\d+){0,3}\s+\S", text)) and len(text) < 60
        if is_heading_style or is_numbered:
            headings.append((style, text))

    print("=" * 78)
    print(f"FILE: {path.name}")
    print(f"HEADINGS: {len(headings)}")
    print("=" * 78)
    for style, text in headings:
        print(f"  [{style}] {text}")

    # collect heading numbers
    nums = set()
    for _, text in headings:
        m = re.match(r"^\s*(\d+(?:\.\d+)*)", text)
        if m:
            nums.add(m.group(1))

    body = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            body += "\n" + " ".join(c.text for c in row.cells)

    refs = set()
    for m in re.finditer(r"(?:详见|参见|见|按照)\s*(\d+(?:\.\d+)*)\s*(?:章节|节)?", body):
        refs.add(m.group(1))
    for m in re.finditer(r"(\d+(?:\.\d+)*)\s*章节", body):
        refs.add(m.group(1))
    for m in re.finditer(r"表\s*(\d+-\d+)", body):
        refs.add(m.group(1))

    print("\n[Cross-reference check]")
    print(f"  heading numbers found ({len(nums)}): "
          f"{', '.join(sorted(nums, key=lambda s: [int(x) for x in s.split('.')]))}")
    unresolved = sorted(r for r in refs if r not in nums)
    if unresolved:
        print(f"  UNRESOLVED references ({len(unresolved)}): {', '.join(unresolved)}")
    else:
        print("  all numeric cross-references resolve to a heading number")
    return 0


if __name__ == "__main__":
    sys.exit(main())
