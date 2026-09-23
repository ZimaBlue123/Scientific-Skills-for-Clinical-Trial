"""PPTX validation utilities."""

from __future__ import annotations
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

def check_overflow(slide) -> bool:
    """Check if text content in a slide overflows its boundaries."""
    # This is a heuristic check as python-pptx doesn't have an exact rendering engine
    # Usually involves checking string length vs placeholder dimensions or auto-fit settings
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        # Naive character limit check for demonstration
        if len(shape.text_frame.text) > 1000:
            return True
    return False

def find_overflow_slides(pptx_path: Path | str) -> list[int]:
    """Find all slide indices that contain text overflow."""
    from pptx import Presentation
    
    path = Path(pptx_path)
    if not path.exists():
        return []
        
    prs = Presentation(str(path))
    overflows = []
    
    for i, slide in enumerate(prs.slides, start=1):
        if check_overflow(slide):
            overflows.append(i)
            
    return overflows

__all__ = ["check_overflow", "find_overflow_slides"]
