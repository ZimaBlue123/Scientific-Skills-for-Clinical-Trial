#!/usr/bin/env python3
"""Dump theme colors, master/layout names, fonts and placeholder geometry of a
reference PPTX so a new deck can be generated with an identical visual identity.

Usage:
    python inspect_pptx_theme_ref.py <pptx> <out.txt> [slide_numbers...]
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu


def inch(v) -> str:
    try:
        return f"{Emu(int(v)).inches:.2f}in"
    except Exception:  # noqa: BLE001
        return str(v)


def fmt_color(font):
    try:
        if font.color is not None and font.color.type is not None:
            if font.color.type == 1:
                return f"theme:{font.color.theme_color}"
            return f"#{font.color.rgb}"
    except Exception:  # noqa: BLE001
        pass
    return "-"


def dump_theme(prs, out):
    out.append("===== THEME (clrScheme) =====")
    try:
        clr = prs.slide_masters[0].element  # placeholder to reach part
        _ = clr
        theme_part = prs.part.package.parts
        for part in theme_part:
            if "theme" in str(part.partname):
                from lxml import etree

                root = etree.fromstring(part.blob)
                ns = {
                    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                }
                for tag in (
                    "dk1",
                    "lt1",
                    "dk2",
                    "lt2",
                    "accent1",
                    "accent2",
                    "accent3",
                    "accent4",
                    "accent5",
                    "accent6",
                    "hlink",
                    "folHlink",
                ):
                    nodes = root.findall(f".//a:clrScheme/a:{tag}/a:srgbClr", namespaces=ns)
                    if not nodes:
                        nodes = root.findall(f".//a:clrScheme/a:{tag}/a:sysClr", namespaces=ns)
                    for n in nodes:
                        val = n.get("val") or n.get("lastClr")
                        out.append(f"  {tag}: #{val}")
                for ftype in ("majorFont", "minorFont"):
                    for n in root.findall(f".//a:fontScheme/a:{ftype}/a:latin", namespaces=ns):
                        out.append(f"  {ftype}/latin: {n.get('typeface')}")
                break
    except Exception as exc:  # noqa: BLE001
        out.append(f"  theme dump failed: {exc}")
    out.append("")


def dump_masters(prs, out):
    out.append("===== MASTERS / LAYOUTS =====")
    for mi, master in enumerate(prs.slide_masters):
        out.append(f"Master {mi}: name={master.name!r}")
        for li, layout in enumerate(master.slide_layouts):
            out.append(f"  Layout {li}: {layout.name!r}")
        # master background fill
        try:
            bg = master.element.find(
                ".//{http://schemas.openxmlformats.org/presentationml/2006/main}bg"
            )
            out.append(f"  master has <p:bg>: {bg is not None}")
        except Exception:  # noqa: BLE001
            pass
    out.append("")


def dump_shapes(slide, out):
    for sh in slide.shapes:
        out.append(
            f"  - {sh.name!r} type={sh.shape_type} "
            f"L={inch(sh.left)} T={inch(sh.top)} W={inch(sh.width)} H={inch(sh.height)}"
        )
        # fill info
        try:
            fill = sh.fill
            out.append(f"      fill.type={fill.type}")
            if fill.type is not None and fill.type == 1:  # solid
                out.append(f"      fill.fg=#{fill.fore_color.rgb}")
        except Exception:  # noqa: BLE001
            pass
        if sh.has_text_frame:
            for pi, para in enumerate(sh.text_frame.paragraphs):
                if not para.text.strip():
                    continue
                out.append(f"      [p{pi}] {para.text[:120]!r} align={para.alignment}")
                for r in para.runs:
                    f = r.font
                    sz = f.size.pt if f.size is not None else None
                    out.append(
                        f"         run font={f.name!r} size={sz} bold={f.bold} color={fmt_color(f)}"
                    )
        if getattr(sh, "has_table", False) and sh.has_table:
            tbl = sh.table
            out.append(f"      TABLE rows={len(tbl.rows)} cols={len(tbl.columns)}")
            widths = [inch(c.width) for c in tbl.columns]
            out.append(f"      colw={widths}")
            for ri, row in enumerate(tbl.rows):
                cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
                out.append(f"      r{ri}: {cells}")
                if ri == 0:
                    for ci, cell in enumerate(row.cells):
                        for para in cell.text_frame.paragraphs:
                            for r in para.runs:
                                f = r.font
                                sz = f.size.pt if f.size is not None else None
                                out.append(
                                    f"        h{ci} font={f.name!r} size={sz} "
                                    f"bold={f.bold} color={fmt_color(f)}"
                                )
                        break


def main(argv: list[str]) -> int:
    src = Path(argv[1])
    out_path = Path(argv[2])
    wanted = [int(x) for x in argv[3:]] if len(argv) > 3 else []

    prs = Presentation(str(src))
    out: list[str] = []
    out.append(f"Slide size: {inch(prs.slide_width)} x {inch(prs.slide_height)}")
    out.append(f"Total slides: {len(prs.slides)}")
    out.append("")

    dump_theme(prs, out)
    dump_masters(prs, out)

    for idx, slide in enumerate(prs.slides, 1):
        if wanted and idx not in wanted:
            continue
        out.append(f"##### SLIDE {idx} layout={slide.slide_layout.name!r} #####")
        dump_shapes(slide, out)
        out.append("")

    out_path.write_text("\n".join(out), encoding="utf-8")
    print(f"wrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
