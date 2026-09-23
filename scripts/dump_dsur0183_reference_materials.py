#!/usr/bin/env python3
"""Dump the DSUR#2 reference materials (xlsx / docx) of review_materials/018-3.

Outputs markdown-ish text files into review_materials/018-3/_work/refs/ so the
auditor can cross-check the DSUR draft against source data.

Usage
-----
    py -3 scripts/dump_dsur0183_reference_materials.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document  # type: ignore
from openpyxl import load_workbook  # type: ignore

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "review_materials" / "018-3" / "前期资料"
OUT = ROOT / "review_materials" / "018-3" / "_work" / "refs"


def dump_xlsx(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=True)
    lines = [f"# XLSX: {path.name}"]
    for ws in wb.worksheets:
        lines.append(f"\n## Sheet: {ws.title}  (dims={ws.dimensions})")
        max_row = min(ws.max_row or 0, 400)
        max_col = min(ws.max_column or 0, 40)
        for r in range(1, max_row + 1):
            vals = []
            for c in range(1, max_col + 1):
                v = ws.cell(row=r, column=c).value
                if v is None:
                    continue
                s = str(v).replace("\n", " ").strip()
                if s:
                    vals.append(f"[{ws.cell(row=r, column=c).coordinate}]{s}")
            if vals:
                lines.append(" | ".join(vals))
    return lines


def dump_docx(path: Path) -> list[str]:
    doc = Document(str(path))
    lines = [f"# DOCX: {path.name}"]
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)
    for i, table in enumerate(doc.tables, 1):
        lines.append(f"\n## Table {i}")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            lines.append(" | ".join(cells))
    return lines


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    for path in sorted(SRC.iterdir()):
        if path.name.startswith("~$"):
            continue
        suf = path.suffix.lower()
        try:
            if suf in {".xlsx", ".xlsm", ".xls"}:
                lines = dump_xlsx(path)
            elif suf == ".docx":
                lines = dump_docx(path)
            else:
                continue
        except Exception as exc:  # noqa: BLE001
            print(f"FAILED {path.name}: {exc}")
            continue
        target = OUT / (path.stem + ".md")
        target.write_text("\n".join(lines), encoding="utf-8")
        print(f"wrote {target.name} ({len(lines)} lines)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
