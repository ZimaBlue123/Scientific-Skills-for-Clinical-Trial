#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Post-delivery self-check for a CTR clinical Excel database.

Run this BEFORE handing the file over. It catches the failures that openpyxl
will happily write out and that only show up later in Excel:
  - zip / workbook integrity
  - blank or duplicated primary keys
  - rows not sorted by the key descending
  - a date written into a "General" cell -> renders as a serial number
  - legacy placeholders ("未登记") left in date columns
  - non-integer sample sizes
  - half-width brackets that were supposed to be normalised
  - auto_filter range not covering the data

Usage:
    python verify_ctr_excel.py <xlsx> [--sheet NAME] [--key 登记号] [--header-row 1]
Exit code 0 = all checks passed.
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
import zipfile

import openpyxl

PLACEHOLDERS = {"未登记", "暂未填写", "暂无", "—", "-", "N/A", "NA"}
DATE_HINTS = ("日期", "公示", "入组", "完成", "启动")
# NB: "主要人群（例）" must NOT be treated as a count column, hence the narrow
# hints below (no bare "例").
COUNT_HINTS = ("样本量", "例数", "入组人数", "受试者人数")


def looks_like_date_col(name: str) -> bool:
    return any(k in name for k in DATE_HINTS)


def looks_like_count_col(name: str) -> bool:
    return any(k in name for k in COUNT_HINTS)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--key", default="登记号")
    ap.add_argument("--header-row", type=int, default=1)
    a = ap.parse_args()

    fails, warns = [], []
    print("=" * 62)
    print("VERIFY:", a.xlsx)
    print("=" * 62)

    # PART 1 - package integrity -------------------------------------------
    try:
        z = zipfile.ZipFile(a.xlsx)
        bad = z.testzip()
        if bad:
            fails.append("PART1 zip corrupt: %s" % bad)
        else:
            print("PART 1  zip integrity .................. PASS")
    except Exception as e:
        print("PART 1  zip integrity .................. FAIL (%s)" % e)
        return 1

    wb = openpyxl.load_workbook(a.xlsx)
    ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
    headers = [ws.cell(a.header_row, c).value for c in range(1, ws.max_column + 1)]
    headers = [h for h in headers if h]
    if a.key not in headers:
        print("PART 2  key column %r missing; headers=%s" % (a.key, headers))
        return 1
    ki = headers.index(a.key)
    last_row = a.header_row
    for r in range(a.header_row + 1, ws.max_row + 1):
        if any(ws.cell(r, c).value is not None for c in range(1, len(headers) + 1)):
            last_row = r
    n = last_row - a.header_row

    # PART 2 - key column ---------------------------------------------------
    keys = [ws.cell(r, ki + 1).value for r in range(a.header_row + 1, last_row + 1)]
    blank = sum(1 for k in keys if k is None or str(k).strip() == "")
    dupes = sorted({k for k in keys if keys.count(k) > 1})
    if blank:
        fails.append("PART2 %d row(s) with blank %s" % (blank, a.key))
    if dupes:
        fails.append("PART2 duplicated %s: %s" % (a.key, dupes))
    print("PART 2  key column (%s) ................ %s  rows=%d blank=%d dup=%d"
          % (a.key, "PASS" if not (blank or dupes) else "FAIL", n, blank, len(dupes)))

    # PART 3 - sort order ---------------------------------------------------
    ordered = all(str(keys[i]) >= str(keys[i + 1]) for i in range(len(keys) - 1))
    if not ordered:
        fails.append("PART3 rows are not sorted by %s descending" % a.key)
    print("PART 3  descending sort ................ %s" % ("PASS" if ordered else "FAIL"))

    # PART 4 - date columns -------------------------------------------------
    dprob, dwarn = [], []
    for c, name in enumerate(headers, 1):
        if not looks_like_date_col(name):
            continue
        for r in range(a.header_row + 1, last_row + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if v is None:
                continue
            if isinstance(v, str):
                if v.strip() in PLACEHOLDERS:
                    dprob.append((r, name, "placeholder %r" % v))
                else:
                    # free text in a date column: usually a deliberate manual
                    # status note -> flag for human confirmation, not a failure
                    dwarn.append((r, name, "free text %r (confirm with user)" % v))
            elif isinstance(v, (dt.datetime, dt.date)):
                if cell.number_format == "General":
                    dprob.append((r, name, "date renders as serial (General)"))
            else:
                dprob.append((r, name, "type %s" % type(v).__name__))
    for p in dprob:
        fails.append("PART4 row %d %s: %s" % p)
    for p in dwarn:
        warns.append("PART4 row %d %s: %s" % p)
    print("PART 4  date columns ................... %s  (%d issue, %d to confirm)"
          % ("PASS" if not dprob else "FAIL", len(dprob), len(dwarn)))

    # PART 5 - count columns ------------------------------------------------
    cprob = []
    for c, name in enumerate(headers, 1):
        if not looks_like_count_col(name):
            continue
        for r in range(a.header_row + 1, last_row + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            if isinstance(v, bool) or not isinstance(v, int):
                cprob.append((r, name, repr(v)[:30]))
    for p in cprob:
        fails.append("PART5 row %d %s: non-integer %s" % p)
    print("PART 5  count columns .................. %s  (%d issue)"
          % ("PASS" if not cprob else "FAIL", len(cprob)))

    # PART 6 - bracket / whitespace normalisation ---------------------------
    bprob = []
    for c, name in enumerate(headers, 1):
        if not (("名称" in name) or ("制品" in name) or ("药物" in name)
                or ("申请人" in name) or ("申办" in name)):
            continue
        for r in range(a.header_row + 1, last_row + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            if "(" in v or ")" in v:
                bprob.append((r, name, v[:30]))
            if re.search(r"\s/\s|/\s", v):
                bprob.append((r, name, "space around '/' -> %s" % v[:30]))
    for p in bprob:
        warns.append("PART6 row %d %s: %s" % p)
    print("PART 6  bracket/space normalisation .... %s  (%d warning)"
          % ("PASS" if not bprob else "WARN", len(bprob)))

    # PART 7 - sheet furniture ---------------------------------------------
    expect = "A%d:%s%d" % (a.header_row,
                           openpyxl.utils.get_column_letter(len(headers)), last_row)
    af_ok = (ws.auto_filter.ref or "").upper() == expect.upper()
    if not af_ok:
        warns.append("PART7 auto_filter %r != %r" % (ws.auto_filter.ref, expect))
    print("PART 7  auto_filter range .............. %s" % ("PASS" if af_ok else "WARN"))

    print("-" * 62)
    for f in fails:
        print("FAIL:", f)
    for w in warns:
        print("WARN:", w)
    print("RESULT: %d failure(s), %d warning(s)" % (len(fails), len(warns)))
    return 1 if fails else 0


if __name__ == "__main__":
    raise SystemExit(main())
