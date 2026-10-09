#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把非劣效分析结果渲染成多子表 Excel 交付稿。

Sheet 结构（与 TVAX-009B Ⅰ期 / TVAX-009-002 Ⅱ期交付稿同构）：
  00_说明    结论文案，八节：目的/数据来源/人群/时点/方法/校准/结论/局限性
  01_汇总    全部 人群 × 分析集 × 指标 的判定一览
  比较N_...  每个「人群_指标」一页，PPS/FAS 各一行

CLI:
  python build_noninferiority_workbook.py --spec spec.json --results results.json \
         --out report.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    print("ERROR: 需要 openpyxl（pip install openpyxl）", file=sys.stderr)
    raise

ILLEGAL_SHEET_CHARS = r"[\/\*\?\[\]:\\]"


# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------


def safe_sheet_name(name: str, used: set, limit: int = 31) -> str:
    clean = re.sub(ILLEGAL_SHEET_CHARS, "_", name)
    if len(clean) > limit:
        clean = clean[:limit]
    base, i = clean, 1
    while clean in used:
        suffix = f"_{i}"
        clean = base[: limit - len(suffix)] + suffix
        i += 1
    used.add(clean)
    return clean


def fmt(v, nd=2):
    return "-" if v is None else f"{v:.{nd}f}"


THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_HEAD = PatternFill("solid", fgColor="F2F2F2")


def style_row(ws, row, ncol, fill=None, bold=False, wrap=False, align="center"):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if bold:
            cell.font = Font(bold=True)
        if fill:
            cell.fill = fill


# ---------------------------------------------------------------------------
# 说明页文案组装
# ---------------------------------------------------------------------------


def auto_summary_lines(results, spec):
    rows = results["rows"]
    lines = []
    idx = 0
    for pop in spec["populations"]:
        for kind, title in (("rate", pop["rate_metric"]["name"]),
                            ("gmc", spec.get("gmc_label", "抗-HBs GMC"))):
            idx += 1
            margin = (spec["margins"]["rate_diff"] if kind == "rate"
                      else spec["margins"]["gmc_ratio"])
            metric_rows = [r for r in rows
                           if r["population"] == pop["name"] and r["metric"] == kind]
            if not metric_rows:
                continue
            head = f"{idx}）{pop['name']}｜{title}"
            if kind == "rate":
                head += f"（界值 {fmt(margin)}%）"
            else:
                head += f"（界值 {margin}）"
            lines.append(head)
            for r in metric_rows:
                if kind == "rate":
                    lines.append(f"   {r['analysis_set']}："
                                 f"{r['trial']['x']}/{r['trial']['n']}="
                                 f"{fmt(r['trial']['pct'])}% vs "
                                 f"{r['control']['x']}/{r['control']['n']}="
                                 f"{fmt(r['control']['pct'])}%，"
                                 f"率差 {fmt(r['estimate']['diff'])}%，"
                                 f"95%CI（{fmt(r['estimate']['lo'])}, "
                                 f"{fmt(r['estimate']['hi'])}）→ {r['verdict']}")
                else:
                    lines.append(f"   {r['analysis_set']}："
                                 f"{fmt(r['trial']['gmc'], 3)} vs "
                                 f"{fmt(r['control']['gmc'], 3)}，"
                                 f"比值 {fmt(r['estimate']['ratio'], 3)}，"
                                 f"95%CI（{fmt(r['estimate']['lo'], 3)}, "
                                 f"{fmt(r['estimate']['hi'], 3)}）→ {r['verdict']}")
    lines.append("   详见第七节的关键解读；【不成立】仅表示现有数据未能确认非劣效，"
                 "并不等同于证实劣效。")
    return lines


def auto_method_lines(spec):
    m = spec["margins"]
    lines = [
        "1）单组率 95%CI：Clopper-Pearson 法。",
        "2）组间率差及其 95%CI：Miettinen-Nurminen 法（含 (n1+n2)/(n1+n2−1) 小样本校正），"
        "并用 Newcombe 杂交 Wilson 法交叉验证。",
        "3）组间率检验：Fisher 确切概率法 / 卡方检验。",
        "4）GMC 及其 95%CI：抗体浓度作对数转换后按 t 分布估计。",
        "5）GMC 比值及其 95%CI：对数转换后成组 t 检验"
        "（与Ⅲ研究 CSR／统计分析报告组间 GMC 比较同口径）；比值 = 试验组 GMC / 对照组 GMC。",
    ]
    cens = spec.get("censoring", {})
    if cens:
        lines.append(f"6）删失值处理：低于检测下限者按 "
                     f"{cens.get('rule_label', cens.get('rule', 'lod_half'))} 参与几何均数计算"
                     f"；该规则须由官方已发表的基线 GMC 反算校准后确定，不得凭经验选取。")
    lines.append(f"7）非劣效界值与判定规则：GMC 比值界值 {m['gmc_ratio']}，"
                 f"95%CI 下限 > {m['gmc_ratio']} 判定成立；"
                 f"率差界值 {m['rate_diff']}%，95%CI 下限 > {m['rate_diff']}% 判定成立。")
    lines.append("8）【不成立】仅表示现有数据未能确认非劣效，并不等同于证实劣效。")
    lines.append("9）数值修约：率、率差及可信区间保留 2 位小数；"
                 "GMC、GMC 比值及可信区间保留 3 位小数。")
    return lines


def auto_calibration_lines(spec):
    lines = ["1）单组 GMC 与 95%CI：应对全部「人群 × 分析集 × 组别」组合逐项吻合。",
             "2）组间 GMC 比较 P 值：应对全部组合逐项吻合。",
             "3）单组率与 n/N：应对全部组合逐项吻合。",
             "4）上述逐项吻合是本表统计口径与官方一致性的直接证据；"
             "任一项不吻合即停止出数，先排查以下四类原因："]
    lines.append("   ① 分析集成员不对：核对排除清单与 preset 分析集名称是否一一对应；")
    lines.append("   ② 删失值规则不对：用官方基线 GMC 反算确认 LOD 取法；")
    lines.append("   ③ FAS 是否走 LOCF：官方汇总表若对缺失值作末次观测结转，"
                 "则 FAS 不得重算，须直接引用官方值；")
    lines.append("   ④ 时对点错位：确认不同免疫程序下「全免后 N 个月」各自落在哪个访视上。")
    return lines


def auto_conclusion_lines(results, spec):
    rows = results["rows"]
    margin_rate = float(spec["margins"]["rate_diff"])
    margin_gmc = float(spec["margins"]["gmc_ratio"])
    rate_rows = [r for r in rows if r["metric"] == "rate"]
    gmc_rows = [r for r in rows if r["metric"] == "gmc"]
    n_fail = sum(1 for r in rows if r["verdict"] == "不成立")

    lines = []
    if n_fail == 0:
        lines.append(f"1）结论：共完成 {len(rows)} 项非劣效比较，"
                     f"各项 95%CI 下限均越过各自界值。")
    elif n_fail == len(rows):
        lines.append(f"1）结论：共完成 {len(rows)} 项非劣效比较，"
                     f"各项 95%CI 下限均未越过各自界值，本次数据未能获得统计学确证。")
    else:
        lines.append(f"1）结论：共完成 {len(rows)} 项非劣效比较，其中 {n_fail} 项的 "
                     f"95%CI 下限未越过各自界值，其余 {len(rows) - n_fail} 项成立。")

    diffs = [r["estimate"]["diff"] for r in rate_rows if r["estimate"]["diff"] is not None]
    if diffs:
        lines.append(f"   ① 率差点估计范围 {fmt(min(diffs))}% ~ {fmt(max(diffs))}%，"
                     f"请结合界值 {fmt(margin_rate)}% 判断其临床意义："
                     f"若点估计量级远小于界值尺度，说明未见实质下降。")
    ratios = [r["estimate"]["ratio"] for r in gmc_rows if r["estimate"]["ratio"] is not None]
    if ratios:
        lines.append(f"   ② GMC 比值点估计范围 {fmt(min(ratios), 3)} ~ {fmt(max(ratios), 3)}，"
                     f"界值 {margin_gmc}。点估计高于界值不代表已确证非劣效，"
                     f"须以 95%CI 下限为准。")

    wides = [(r, r["estimate"]["hi"] / r["estimate"]["lo"])
             for r in gmc_rows
             if r["estimate"]["lo"] and r["estimate"]["lo"] > 0]
    if wides:
        widest = max(wides, key=lambda x: x[1])
        lines.append(f"   ③ 未获确证的常见原因是精度不足："
                     f"本组 GMC 比值 95%CI 上下限之比最大达 {widest[1]:.1f} 倍"
                     f"（{widest[0]['population']} {widest[0]['analysis_set']}），"
                     f"提示样本量小或个体变异大。")

    ceil = [r for r in rate_rows if r["control"]["pct"] is not None
            and r["control"]["pct"] >= 100.0]
    if ceil:
        names = "、".join(f"{r['population']} {r['analysis_set']}" for r in ceil)
        lines.append(f"   ④ 天花板效应：{names} 的对照组率为 100.00%，"
                     f"会系统性拉大负向率差并压缩非劣效成立空间，解读时应注明。")
    return lines


def build_notes_lines(spec, results):
    narrative = spec.get("narrative", {})
    lines = [spec.get("notes_title", "疫苗免疫原性非劣效分析 —— 说明页"), ""]
    sections = [
        ("【核心结论速览】", auto_summary_lines(results, spec)),
        ("【一、分析目的】", narrative.get("purpose", [])),
        ("【二、数据来源】", narrative.get("data_source", [])),
        ("【三、人群与组别】", narrative.get("arm_definitions", [])),
        ("【四、时点口径】", narrative.get("timepoint_note", [])),
        ("【五、统计方法与非劣效判定】", auto_method_lines(spec)),
        ("【六、与官方发表值的校准】",
         auto_calibration_lines(spec) + narrative.get("calibration_notes", [])),
        ("【七、最终结论与依据】",
         auto_conclusion_lines(results, spec) + narrative.get("interpretation", [])),
        ("【八、局限性说明】", narrative.get("limitations", [])),
    ]
    for title, body in sections:
        if title:
            lines.append(title)
        for item in body:
            lines.append(item)
        lines.append("")
    return lines


# ---------------------------------------------------------------------------
# 写 workbook
# ---------------------------------------------------------------------------


def write_notes(ws, lines):
    ws.column_dimensions["A"].width = 150
    for i, line in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if i == 1:
            c.font = Font(bold=True, size=13)


N_SUMMARY_COLS = 10


def write_summary(ws, results, spec):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=N_SUMMARY_COLS)
    title = (f"非劣效判定汇总（{spec.get('summary_title', '')}"
             f"｜GMC 比值界值 {spec['margins']['gmc_ratio']}；"
             f"率差界值 {spec['margins']['rate_diff']}%）")
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)

    heads = ["人群", "指标", "分析集", "试验组", "对照组",
             "估计量（率差%／GMC比值）", "95%CI下限", "95%CI上限", "P值", "非劣效判定"]
    for i, h in enumerate(heads, start=1):
        ws.cell(row=2, column=i, value=h)
    style_row(ws, 2, N_SUMMARY_COLS, fill=FILL_HEAD, bold=True, wrap=True)

    r = 3
    for row in results["rows"]:
        if row["metric"] == "rate":
            t = f"{row['trial']['x']}/{row['trial']['n']}（{fmt(row['trial']['pct'])}%）"
            c = f"{row['control']['x']}/{row['control']['n']}（{fmt(row['control']['pct'])}%）"
            est, lo, hi, nd = (row["estimate"]["diff"], row["estimate"]["lo"],
                               row["estimate"]["hi"], 2)
        else:
            t = f"N={row['trial']['n']}，GMC {fmt(row['trial']['gmc'], 3)}"
            c = f"N={row['control']['n']}，GMC {fmt(row['control']['gmc'], 3)}"
            est, lo, hi, nd = (row["estimate"]["ratio"], row["estimate"]["lo"],
                               row["estimate"]["hi"], 3)
        vals = [row["population"], row["metric_label"], row["analysis_set"], t, c,
                round(est, nd) if est is not None else None,
                round(lo, nd) if lo is not None else None,
                round(hi, nd) if hi is not None else None,
                round(row["estimate"]["p"], 4), row["verdict"]]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)
        style_row(ws, r, N_SUMMARY_COLS)
        r += 1

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 24
    for col in ("C", "D", "E", "F"):
        ws.column_dimensions[col].width = 20
    for col in ("G", "H", "I", "J"):
        ws.column_dimensions[col].width = 12


N_COMPARE_COLS = 13


def write_compare_sheet(ws, comp_rows, spec, kind):
    pop_label = comp_rows[0]["population"]
    metric_label = comp_rows[0]["metric_label"]
    margin = comp_rows[0]["margin"]
    kind_cn = "率差" if kind == "rate" else "GMC比值"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=N_COMPARE_COLS)
    ws.cell(row=1, column=1,
            value=f"{pop_label}｜{metric_label}（{spec['arms']['trial']} vs "
                  f"{spec['arms']['control']}，{spec.get('summary_title', '')}）")
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=N_COMPARE_COLS)
    ws.cell(row=2, column=1,
            value=f"人群：{pop_label}｜指标：{metric_label}｜"
                  f"非劣效界值：{fmt(margin, 2 if kind == 'rate' else 2)}"
                  f"{'%' if kind == 'rate' else ''}")
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="left", vertical="center",
                                                   wrap_text=True)

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=N_COMPARE_COLS)
    ws.cell(row=3, column=1, value=spec.get(
        "compare_desc", "试验组相对于对照组的非劣效比较；详见 00_说明 页。"))
    ws.cell(row=3, column=1).alignment = Alignment(horizontal="left", vertical="center",
                                                   wrap_text=True)

    ws.merge_cells(start_row=7, start_column=2, end_row=7, end_column=5)
    ws.cell(row=7, column=2, value="试验组（拟证明非劣）")
    ws.merge_cells(start_row=7, start_column=6, end_row=7, end_column=9)
    ws.cell(row=7, column=6, value="对照组")
    ws.merge_cells(start_row=7, start_column=10, end_row=7, end_column=13)
    ws.cell(row=7, column=10, value="组间比较")
    style_row(ws, 7, N_COMPARE_COLS, fill=FILL_HEAD, bold=True)

    if kind == "rate":
        heads = ["分析集", "n/N", "率(%)", "95%CI下限", "95%CI上限",
                 "n/N", "率(%)", "95%CI下限", "95%CI上限",
                 "率差(%)", "95%CI下限", "95%CI上限", "非劣效判定"]
    else:
        heads = ["分析集", "N", "GMC", "95%CI下限", "95%CI上限",
                 "N", "GMC", "95%CI下限", "95%CI上限",
                 "GMC比值", "95%CI下限", "95%CI上限", "非劣效判定"]
    for i, h in enumerate(heads, start=1):
        ws.cell(row=8, column=i, value=h)
    style_row(ws, 8, N_COMPARE_COLS, fill=FILL_HEAD, bold=True, wrap=True)

    nd = 2 if kind == "rate" else 3
    r = 9
    for row in comp_rows:
        if kind == "rate":
            vals = [row["analysis_set"],
                    f"{row['trial']['x']}/{row['trial']['n']}",
                    row["trial"]["pct"], row["trial"]["lo"], row["trial"]["hi"],
                    f"{row['control']['x']}/{row['control']['n']}",
                    row["control"]["pct"], row["control"]["lo"], row["control"]["hi"],
                    row["estimate"]["diff"], row["estimate"]["lo"], row["estimate"]["hi"],
                    row["verdict"]]
        else:
            vals = [row["analysis_set"], row["trial"]["n"],
                    row["trial"]["gmc"], row["trial"]["lo"], row["trial"]["hi"],
                    row["control"]["n"],
                    row["control"]["gmc"], row["control"]["lo"], row["control"]["hi"],
                    row["estimate"]["ratio"], row["estimate"]["lo"], row["estimate"]["hi"],
                    row["verdict"]]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i,
                    value=round(v, nd) if isinstance(v, float) else v)
        style_row(ws, r, N_COMPARE_COLS)
        r += 1

    note = r + 1
    ws.cell(row=note, column=1, value="判定规则：")
    ws.merge_cells(start_row=note, start_column=2, end_row=note, end_column=N_COMPARE_COLS)
    if kind == "rate":
        rule = (f"{kind_cn}（试验组 − 对照组）的 95%CI 下限 > {fmt(margin)}% 判定非劣效成立；"
                f"否则为不成立。【不成立】仅表示现有数据未能确认非劣效，并不等同于证实劣效。"
                f"率差与 95%CI 用 Miettinen-Nurminen 法（含小样本校正），"
                f"单组率 95%CI 用 Clopper-Pearson 法，组间检验用 Fisher 确切概率法。")
    else:
        rule = (f"{kind_cn}（试验组/对照组）的 95%CI 下限 > {margin} 判定非劣效成立；"
                f"否则为不成立。【不成立】仅表示现有数据未能确认非劣效，并不等同于证实劣效。"
                f"GMC 比值为抗体浓度对数转换后成组 t 检验所得。")
    extra = spec.get("compare_rule_extra", "")
    ws.cell(row=note, column=2, value=rule + (" " + extra if extra else ""))
    ws.cell(row=note, column=2).alignment = Alignment(horizontal="left", vertical="center",
                                                      wrap_text=True)
    ws.row_dimensions[note].height = 46

    ws.column_dimensions["A"].width = 10
    for c in range(2, N_COMPARE_COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="渲染非劣效分析 Excel 交付稿")
    ap.add_argument("--spec", required=True)
    ap.add_argument("--results", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args(argv)

    with open(args.spec, encoding="utf-8") as f:
        spec = json.load(f)
    with open(args.results, encoding="utf-8") as f:
        results = json.load(f)

    wb = Workbook()
    notes = wb.active
    notes.title = "00_说明"
    write_notes(notes, build_notes_lines(spec, results))

    write_summary(wb.create_sheet("01_汇总"), results, spec)

    used = {"00_说明", "01_汇总"}
    groups = []
    for row in results["rows"]:
        key = (row["population"], row["metric"])
        if key not in [(g[0], g[1]) for g in groups]:
            groups.append((row["population"], row["metric"], row["metric_label"]))
    for i, (pop, kind, label) in enumerate(groups, start=1):
        comp_rows = [r for r in results["rows"]
                     if r["population"] == pop and r["metric"] == kind]
        prefix = f"比较{i}_"
        name = safe_sheet_name(prefix + f"{pop}_{label}", used)
        write_compare_sheet(wb.create_sheet(name), comp_rows, spec, kind)

    out_dir = os.path.dirname(os.path.abspath(args.out))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(args.out)
    print(f"已写出：{args.out}")
    print("sheets：" + "，".join(wb.sheetnames))
    return 0


if __name__ == "__main__":
    sys.exit(main())
