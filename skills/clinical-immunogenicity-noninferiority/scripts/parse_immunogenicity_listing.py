#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从统计分析报告（.docx）抽取个体级免疫原性清单。

为什么不能按表头关键词找表：同一册里往往有多张表头高度相似的表。
TVAX-009B Ⅰ期统计报告第 2 册中，「表16.2.5.9 乙肝两对半检测清单」与
「表16.2.6.1 抗HBs结果清单」的表头都同时含 研究编号 / HBsAb / HBcAb，
用关键词匹配会命中前者（多一列「采样时间」），解析结果全错。
正确做法：遍历 body，遇到题注段落（如「表16.2.6.1 抗HBs结果清单(FAS)」）
之后取紧接着的第一个 w:tbl。

CLI:
  python parse_immunogenicity_listing.py --spec spec.json --out subjects.csv
  python parse_immunogenicity_listing.py --spec spec.json --dry-run
        # 只打印定位到的表头与列映射，不写文件（排查用）
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)
from compute_noninferiority import parse_conc  # noqa: E402

try:
    import docx
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    print("ERROR: 需要 python-docx（pip install python-docx）", file=sys.stderr)
    raise


# ---------------------------------------------------------------------------
# docx 辅助
# ---------------------------------------------------------------------------


def row_cells_unique(row):
    """python-docx 对合并单元格会重复返回 tc；按元素身份去重。"""
    out, seen = [], set()
    for tc in row.cells:
        key = id(tc._tc)
        if key in seen:
            continue
        seen.add(key)
        out.append(" ".join(tc.text.split()))
    return out


def locate_table_by_caption(doc, caption_prefix):
    """按题注定位表格：题注段落之后的第一个 w:tbl。"""
    n_before = 0
    pending = False
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            if pending:
                return doc.tables[n_before]
            n_before += 1
        elif child.tag == qn("w:p"):
            text = "".join(n.text or "" for n in child.iter(qn("w:t"))).strip()
            if text.startswith(caption_prefix):
                pending = True
    return None


def header_index_map(header_cells, key, occurrence=1):
    """返回第 occurrence 个以 key 开头的列下标（occurrence 从 1 计）。找不到返回 None。"""
    hits = [i for i, h in enumerate(header_cells) if h.startswith(key)]
    if len(hits) < occurrence:
        return None
    return hits[occurrence - 1]


# ---------------------------------------------------------------------------
# 解析
# ---------------------------------------------------------------------------


def parse_exclusions(source: dict) -> dict:
    """返回 {受试者编号: {分析集: 排除原因}}（仅保留有原因的项）。"""
    path = source.get("exclusions_docx")
    caption = source.get("exclusions_caption")
    if not path or not caption:
        return {}
    doc = docx.Document(path)
    table = locate_table_by_caption(doc, caption)
    if table is None:
        raise RuntimeError(f"未定位到清单表：{caption} @ {path}")
    head = row_cells_unique(table.rows[0])
    id_col = header_index_map(head, source.get("id_col", "研究编号"))

    set_cols = {}
    for a_set in source.get("analysis_sets", []):
        idx = header_index_map(head, f"未进入{a_set}")
        if idx is not None:
            set_cols[a_set] = idx

    pattern = re.compile(source.get("id_pattern", r"^[A-Za-z]\d+$"))
    out = {}
    for row in table.rows[1:]:
        cells = row_cells_unique(row)
        sid = cells[id_col].strip()
        if not pattern.match(sid):
            continue
        reasons = {a_set: cells[idx].strip() for a_set, idx in set_cols.items()}
        if any(reasons.values()):
            out[sid] = reasons
    return out


def parse_listing(source: dict, exclusions: dict) -> tuple[list[dict], list[str], dict]:
    """返回 (记录, 表头, 列映射)。"""
    doc = docx.Document(source["listing_docx"])
    table = locate_table_by_caption(doc, source["listing_caption"])
    if table is None:
        raise RuntimeError(f"未定位到清单表：{source['listing_caption']}")

    h_row = int(source.get("header_row", 1))
    header = row_cells_unique(table.rows[h_row])

    id_col = source.get("id_col", "研究编号")
    arm_col = source.get("arm_col", "组别")
    colmap = {
        "id": header_index_map(header, id_col),
        "arm": header_index_map(header, arm_col),
    }
    for col in source.get("static_cols", []):
        colmap[col] = header_index_map(header, col)
    for m in source["measures"]:
        key = f"field::{m['field']}"
        colmap[key] = header_index_map(header, m.get("header", m["field"]),
                                       int(m.get("occurrence", 1)))
    missing = [k for k, v in colmap.items() if v is None]
    if missing:
        raise RuntimeError(f"表头未找到列：{missing}\n实际表头：{header}")

    pattern = re.compile(source.get("id_pattern", r"^[A-Za-z]\d+$"))
    out = []
    for row in table.rows[h_row + 1:]:
        cells = row_cells_unique(row)
        sid = cells[colmap["id"]].strip()
        if not pattern.match(sid):
            continue
        rec = {"sid": sid, "group": cells[colmap["arm"]].strip()}
        for col in source.get("static_cols", []):
            rec[col] = cells[colmap[col]].strip()
        for a_set in source.get("analysis_sets", []):
            rec[f"in_{a_set.lower()}"] = not bool(exclusions.get(sid, {}).get(a_set))
        for m in source["measures"]:
            field = m["field"]
            raw_col = colmap["field::" + field]
            raw = cells[raw_col].strip()
            value, flag = parse_conc(raw, tuple(source.get("below_tokens", ("＜", "<"))),
                                     tuple(source.get("above_tokens", ("＞", ">"))))
            rec[field + "_raw"] = raw
            rec[field + "_value"] = "" if value is None else value
            rec[field + "_flag"] = flag
        out.append(rec)
    return out, header, colmap


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="抽取个体级免疫原性清单")
    ap.add_argument("--spec", required=True, help="分析规格 JSON")
    ap.add_argument("--out", help="输出 CSV 路径")
    ap.add_argument("--dry-run", action="store_true", help="只打印定位结果")
    args = ap.parse_args(argv)

    with open(args.spec, encoding="utf-8") as f:
        spec = json.load(f)
    source = spec["source"]

    exclusions = parse_exclusions(source)
    records, header, colmap = parse_listing(source, exclusions)

    # 列 <=> 字段映射：便于人工核对occurrence是否取对
    field_desc = {k: (v, header[v]) for k, v in colmap.items() if v is not None}

    if args.dry_run:
        print(f"清单表头（第 {source.get('header_row', 1)} 行）：")
        print("  " + " | ".join(header))
        print("列映射：")
        for name, (idx, head) in field_desc.items():
            print(f"  {name:<16} -> 第 {idx} 列 「{head}」")
        print(f"解析记录：{len(records)} 例；被排除：{len(exclusions)} 例")
        for sid, reasons in sorted(exclusions.items()):
            print(f"  {sid}: " + "; ".join(f"{k}={v}" for k, v in reasons.items() if v))
        return 0

    if not args.out:
        ap.error("需要 --out 或 --dry-run")

    fields = ["sid", "group", *source.get("static_cols", [])]
    fields += [f"in_{s.lower()}" for s in source.get("analysis_sets", [])]
    for m in source["measures"]:
        fields += [m["field"] + "_raw", m["field"] + "_value", m["field"] + "_flag"]

    os.makedirs(os.path.dirname(os.path.abspath(args.out)) or ".", exist_ok=True)
    with open(args.out, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for r in records:
            writer.writerow(r)

    arm_count = {}
    for r in records:
        arm_count[r["group"]] = arm_count.get(r["group"], 0) + 1
    print(f"已写出：{args.out}（{len(records)} 例）")
    print("组别分布：" + "，".join(f"{k} {v} 例" for k, v in arm_count.items()))
    if not records:
        print("WARN: 0 条记录，请检查 header_row / id_pattern / occurrence 配置",
              file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
